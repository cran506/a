"""由本次仿真数据生成中文 Markdown 与可提交的 Word 实验报告。"""
from pathlib import Path
import math

import numpy as np


def _number(value, digits=4):
    """None 表示未观测到事件；不把缺失事件替换为零。"""
    if value is None:
        return "—"
    try:
        value = float(value)
    except (ValueError, TypeError):
        return str(value)
    if not math.isfinite(value):
        return "—"
    return f"{value:.{digits}f}"


def _relative_change(baseline, other):
    if baseline is None or other is None or baseline == 0:
        return "无法计算相对变化"
    change = 100. * (float(other) / float(baseline) - 1.)
    direction = "增大" if change >= 0 else "减小"
    return f"{direction} {_number(abs(change), 2)}%"


def _table(headers, rows):
    return {"type": "table", "headers": headers, "rows": rows}


def _paragraph(text):
    return {"type": "paragraph", "text": text}


def _heading(text, level=1):
    return {"type": "heading", "text": text, "level": level}


def _picture(filename, caption):
    return {"type": "picture", "filename": filename, "caption": caption}


def _metric_rows(results, names, specifications):
    return [[label] + [_number(results[name]["metrics"].get(key), digits)
                      for name in names]
            for label, key, digits in specifications]


def _homework_blocks(results, suffix, label, filename):
    names = [f"positional_{suffix}", f"incremental_{suffix}"]
    position, increment = (results[name] for name in names)
    pm, im = position["metrics"], increment["metrics"]
    alpha = position["config"]["alpha"]
    blocks = [_heading(f"{label}（α = {alpha:g}）", 2)]
    blocks.append(_table(["指标", "位置式 PID", "增量式 PID"],
        _metric_rows(results, names, [
            ("真实误差首次 ≤ 0 / s", "first_true_error_nonpositive_s", 4),
            ("量化反馈误差首次 ≤ 0 / s", "first_sampled_error_nonpositive_s", 4),
            ("首次退出 +10 V / s", "first_exit_positive_saturation_s", 4),
            ("首次退出 − 量化误差过零 / s", "first_exit_minus_error_zero_s", 4),
            ("量化误差过零后退出 +10 V / s", "exit_positive_saturation_after_zero_s", 4),
            ("量化误差过零时输出 / V", "voltage_at_sampled_zero_v", 4),
            ("超调 / mm", "overshoot_mm", 4),
            ("实际位移峰值 / mm", "peak_position_mm", 4),
            ("±10 V 饱和总时长 / s", "saturation_time_s", 4),
            ("限幅前候选电压最大值 / V", "max_internal_voltage_v", 4),
            ("限幅前候选电压最小值 / V", "min_internal_voltage_v", 4),
            ("末 200 ms 真实误差 RMS / mm", "final_200ms_rms_error_mm", 4),
        ])))
    initial = [_number(results[name]["data"]["raw_voltage_v"][0], 2)
               for name in names]
    blocks.append(_paragraph(
        f"两组首采样限幅前电压分别为 {initial[0]} V、{initial[1]} V，"
        "实际输出均受 ±10 V 约束。阶跃在 t=0 输入，控制器以零误差历史初始化，"
        "所以微分项包含参考阶跃引起的 derivative kick。退饱和事件由实际采样序列提取；"
        "首次退出可以早于误差过零，负的时间差必须保留，不能用零代替。"))
    blocks.append(_paragraph(
        f"本组位置式超调 {_number(pm.get('overshoot_mm'))} mm，"
        f"增量式超调 {_number(im.get('overshoot_mm'))} mm；"
        f"增量式相对位置式{_relative_change(pm.get('overshoot_mm'), im.get('overshoot_mm'))}。"
        f"两组首次退出 +10 V 的时刻为 {_number(pm.get('first_exit_positive_saturation_s'))} s、"
        f"{_number(im.get('first_exit_positive_saturation_s'))} s，"
        f"量化误差过零时输出为 {_number(pm.get('voltage_at_sampled_zero_v'))} V、"
        f"{_number(im.get('voltage_at_sampled_zero_v'))} V。"
        "应结合图中的积分状态、限幅前候选量和实际电压判断迟滞，"
        "不能只用一次退饱和时间代表全过程。"))
    blocks.append(_picture(filename,
        f"{label}：位移、误差、实际/候选控制电压与控制器记忆状态。"))
    return blocks


def _build_report(results, probe_metrics):
    required = ("conservative", "aggressive", "unfiltered", "positional_raw",
                "incremental_raw", "positional_filtered", "incremental_filtered")
    missing = [name for name in required if name not in results]
    if missing:
        raise ValueError("报告缺少实验结果：" + ", ".join(missing))
    a, b, raw = (results[name] for name in required[:3])
    ca, cb, cr = (item["config"] for item in (a, b, raw))
    ma, mb = a["metrics"], b["metrics"]
    scurve_names = list(required[:3])
    t = a["data"]["time_s"]
    ts = ca["ts"]
    blocks = [
        _heading("第五章 数字 PID 闭环控制实验报告", 0),
        _paragraph("基于第三、四、五章课件、上传 HydraulicTwin 与 Word 中的数字 PID、"
                   "S 型轨迹代码，完成三组课堂实验及教材习题 5-1。"
                   "报告内数值来自本次非线性闭环运行；示意图中的 65% 增益衰减、"
                   "约 17.4 mm 滞后和稳定性描述不作为预设输出。"),
        _heading("1 实验目标与依据"),
        _paragraph("课堂实验分别比较保守增益、激进增益和未滤波微分，"
                   "观察 S 型跟踪、量化噪声及变负载下流量增益与动态误差的变化。"
                   "习题 5-1 在相同增益和限幅下比较位置式历史积分与增量式截断输出记忆，"
                   "提取误差首次非正和退出 +10 V 的时刻。"
                   "课件的 1.0 s 完成 100 mm 定位、误差小于 0.05 mm 是验收目标，"
                   "是否达到须由实际误差判断。"),
        _heading("2 模型、采样与参数"),
        _paragraph("上传的 hydraulic_digital_twin.py 保持原样。"
                   "闭环子类按第三章孔口方程补充压差符号回流，避免超过供油或回油压力时"
                   "仍强制原流向；保留原四状态液压动力学，并增加一阶实际阀芯响应。"
                   "低压体积模量采用蒸气压 Pv=2 kPa 的气穴等效近似，"
                   "不构成完整多相气穴模型。压力不通过事后裁剪伪造。"
                   "原 Twin 在 x=0.4 m 激活的工件接触负载在闭环中禁用，"
                   "替换为明确的时间窗外载，避免两个负载叠加。"),
        _table(["项目", "本次设定与含义"], [
            ["液压主体", "质量 500 kg；A₁=0.010 m²、A₂=0.005 m²；供油 21 MPa、回油 0.1 MPa；其余见 physical_config.json"],
            ["初始状态", f"初始绝对位置 {_number(ca.get('origin_m'), 6)} m 对齐 ADC 代码；速度为零，两腔初压 0.1/0.2 MPa 实现液压力平衡；报告位移相对该原点"],
            ["采样及执行", f"Ts={ts * 1000:g} ms，控制电压 ±10 V，阀命令为 u/10；每周期零阶保持（ZOH）"],
            ["传感器", f"{ca['adc_bits']} bit、0～{ca['sensor_span'] * 1000:g} mm；LSB={_number(ca['sensor_lsb_mm'], 8)} mm；四舍五入量化，代码饱和 0～4095"],
            ["S 型参考", "目标 100 mm，运动段 1.0 s，总时长 1.4 s；七段比例 0.10/0.15/0.10/0.30/0.10/0.15/0.10"],
            ["工艺外载", f"t∈[0.35,0.70) s 施加 {_number(ca['load_force'] / 1000, 1)} kN 阻力，其余时刻为零；加载/卸载与采样边界对齐"],
            ["阀芯滞后", f"τ={ca['valve_tau'] * 1000:g} ms；补充一阶响应参数"],
            ["数值积分", f"每个 ZOH 周期调用隐式 Radau；相对容差 rtol={ca['rtol']:g}；切换时窗逐周期锁存"],
        ]),
        _paragraph(f"{ca['load_force']/1000:g} kN 人为外载、{ca['valve_tau']*1000:g} ms 阀芯时间常数、Pv=2 kPa 和 physical_config.json 中"
                   "补齐的参数是用于复现教学闭环的明确仿真假设。"
                   "这些参数未经真实设备辨识；程序运行和数值检查通过并不等于模型已经得到实验验证。"
                   "因此不能把本次曲线直接解释为实际液压机的已验证性能。"),
        _paragraph("原对象未配置补油阀、溢流阀及管路分布参数。卸载时腔压可能高于 21 MPa "
                   "供油压力，低压腔也可能接近 2 kPa 气穴近似下限；"
                   "本报告按实际压力极值原样记录。约 30 MPa 的瞬态峰值属于这一简化"
                   "封闭模型的响应，不能当作实际系统已得到安全验证的运行压力。"
                   "在真实设备分析中需要补齐上述压力保护、补油及气穴物理。"),
        _paragraph(f"感知→量化→PID→限幅→ZOH→Radau 的顺序每 {_number(ts * 1000, 0)} ms 执行一次。"
                   f"时间刻度为 t[k]=kTs，共 {len(t)} 个采样点，包含两端；"
                   f"{_number(t[-1], 3)} s 对应 N={len(t)-1} 个物理积分区间，CSV 记录 N+1 行。"
                   "末端电压是该采样点的决策值，不计入已经完成的 ZOH 饱和持续时间。"),
        _heading("3 轨迹与数字 PID 实现"),
        _paragraph("七段 S 型轨迹逐段对常跃度解析积分，联合生成位置、速度和加速度。"
                   "位置、速度、加速度在段边界连续，跃度有界并允许阶跃；"
                   "t≥1.0 s 后参考保持 100 mm，速度和加速度为零。"
                   "解析边界条件消除原示例 linspace 点数与 dt 不一致、逐级数值积分累积偏差的问题。"),
        _paragraph("Twin 内部采用 SI 单位。PID 接口接收 m，在内部将误差乘 100 转为 cm；"
                   "Kp 单位为 V/cm，Ki 为 V/(cm·s)，Kd 为 V·s/cm，输出为 V。"
                   "换成毫米或米时必须同步换算增益，否则无法复现 Word 的整定。"),
        _paragraph("不完全微分：D[k]=αD[k−1]+(1−α)(Kd/Ts)(e[k]−e[k−1])。"
                   "增量式：Δu[k]=Kp(e[k]−e[k−1])+Ki e[k]Ts+D[k]−D[k−1]；"
                   "u[k]=clip(u[k−1]+Δu[k],−10,+10)，截断后的实际 u[k] 写回记忆。"
                   "位置式：I[k]=I[k−1]+Ki e[k]Ts；"
                   "u_calc[k]=Kp e[k]+I[k]+D[k]，只对最终输出限幅，内部 I[k] 继续累加。"),
        _paragraph("课堂组保留 Word 的积分分离：|e|≥0.15 cm（1.5 mm）时"
                   "积分增量乘 0.05，小误差恢复 1。习题 5-1 两类控制器均关闭积分分离，"
                   "防止额外抗饱和措施掩盖位置式与增量式的结构差异。"
                   "增量式的 raw_voltage_v 表示本步限幅前候选值，不是跨周期保留的未截断输出。"),
        _heading("4 三组课堂实验结果"),
        _paragraph("Word 给出课堂 Kp=3.8/16.0 和 α=0.85/0，但没有提供课堂 Ki、Kd 的"
                   "具体实例化值。本程序公开补充默认 Ki=22.0、Kd=0.06，取课件习题中的数值"
                   "作为共同基准；它们不是对上传 Twin 已完成辨识或稳定整定的结果。"),
        _table(["实验", "Kp / V·cm⁻¹", "Ki / V·cm⁻¹·s⁻¹", "Kd / V·s·cm⁻¹", "α"], [
            ["保守增益 + 滤波", _number(ca["kp"], 1), _number(ca["ki"], 1), _number(ca["kd"], 2), _number(ca["alpha"], 2)],
            ["激进增益 + 滤波", _number(cb["kp"], 1), _number(cb["ki"], 1), _number(cb["kd"], 2), _number(cb["alpha"], 2)],
            ["保守增益 + 无滤波", _number(cr["kp"], 1), _number(cr["ki"], 1), _number(cr["kd"], 2), _number(cr["alpha"], 2)],
        ]),
        _table(["真实闭环指标", "保守 + 滤波", "激进 + 滤波", "保守 + 无滤波"],
            _metric_rows(results, scurve_names, [
                ("重载最大绝对误差 / mm", "heavy_peak_abs_error_mm", 4),
                ("重载误差 RMS / mm", "heavy_rms_error_mm", 4),
                ("1.0 s 真实误差 / mm", "error_at_1s_mm", 4),
                ("末 200 ms 最大绝对误差 / mm", "final_200ms_max_abs_error_mm", 4),
                ("末 200 ms 误差 RMS / mm", "final_200ms_rms_error_mm", 4),
                ("末端真实误差 / mm", "final_error_mm", 4),
                ("1.1 s 至末端误差峰峰值 / mm", "unloaded_1p1_to_end_error_p2p_mm", 4),
                ("1.1 s 至末端电压峰峰值 / V", "unloaded_1p1_to_end_voltage_p2p_v", 4),
                ("最大绝对输出电压 / V", "peak_abs_voltage_v", 4),
                ("最低绝对腔压 / MPa", "min_absolute_pressure_mpa", 6),
                ("最高绝对腔压 / MPa", "max_absolute_pressure_mpa", 4),
            ])),
        _picture("classroom_comparison.png", "图 1　三组课堂实验的 S 型位移跟踪、真实误差和控制电压；灰色区域为工艺外载时窗。"),
        _paragraph(f"激进组重载误差 RMS 为 {_number(mb['heavy_rms_error_mm'])} mm，"
                   f"相对保守组的 {_number(ma['heavy_rms_error_mm'])} mm"
                   f"{_relative_change(ma['heavy_rms_error_mm'], mb['heavy_rms_error_mm'])}；"
                   f"卸载后 1.1 s 至末端真实误差峰峰值分别为 "
                   f"{_number(ma['unloaded_1p1_to_end_error_p2p_mm'])} mm 和 "
                   f"{_number(mb['unloaded_1p1_to_end_error_p2p_mm'])} mm。"
                   "重载跟踪与卸载波动须同时比较，单纯提高比例增益不能自动证明整体性能改善。"
                   "这里记录的是有限时窗内的波动，未进行相位裕度或极限环的严格稳定性证明。"),
        _paragraph(f"本次阀芯响应为 {ca['valve_tau']*1000:g} ms，保守组卸载后 1.1 s 至末端"
                   f"实测误差峰峰值为 {_number(ma['unloaded_1p1_to_end_error_p2p_mm'])} mm。"
                   "不能依据较低的 Kp 直接断言“保守组卸载无振荡”。"
                   "真实非线性 Twin 的结果与示意图不同，是本实验需要说明的结果；"
                   "没有人为叠加正弦信号，也没有通过画图强制制造或消除振荡。"),
        _heading("5 流量增益机理与量化噪声"),
        _paragraph("正向伸出且 Ps>P₁ 时，Kq=Cd·w·√[2(Ps−P₁)/ρ]，单位 m²/s。"
                   "负载提高驱动腔压力会缩小阀口有效压差，使 Kq 按平方根规律下降；"
                   "卸载则使物理增益回升。统计轻载区间固定为 [0.25,0.34) s，"
                   "重载区间固定为 [0.50,0.65) s，取各区间中位数以避开加载瞬态。"
                   "下表的下降百分比从两区间 Kq 中位数计算，不使用预设 65%。"),
        _paragraph("全时 Kq 曲线按当前 P₁ 计算伸出支路的潜在增益，未按实际阀芯方向筛选。"
                   "阀反向时，它不代表当前反向支路的局部流量增益；Ps≤P₁ 时记录为缺失值。"
                   "本次默认轻载与重载统计窗内阀芯均为正向，下降百分比用于这两处有效伸出工况。"),
        _table(["机理指标", "保守 + 滤波", "激进 + 滤波", "保守 + 无滤波"],
            _metric_rows(results, scurve_names, [
                ("轻载 P₁ 中位数 / MPa", "light_median_p1_mpa", 4),
                ("重载 P₁ 中位数 / MPa", "heavy_median_p1_mpa", 4),
                ("轻载 Kq 中位数 / m²·s⁻¹", "light_median_kq_m2_s", 6),
                ("重载 Kq 中位数 / m²·s⁻¹", "heavy_median_kq_m2_s", 6),
                ("重载 Kq 下降 / %", "kq_decline_percent", 2),
            ])),
        _picture("hydraulic_mechanism.png", "图 2　由实际驱动腔压力、阀口压差和 Kq 变化展示变负载机理。"),
        _paragraph(f"12 bit、300 mm 量程的 1 LSB 为 {_number(ca['sensor_lsb_mm'], 8)} mm。"
                   "Ts=1 ms 时，后向差分中的 1/Ts=1000 s⁻¹；"
                   f"在 Kd={ca['kd']:g} V·s/cm 下，1 LSB 单步跳变对应 "
                   f"{_number(ca['kd'] * (ca['sensor_lsb_mm'] / 10) / ts, 6)} V 的未滤波微分变化。"
                   "不能把 1000 倍直接当作电压幅值，它还依赖单位换算与 Kd。"),
        _paragraph("闭环未滤波组与滤波组的轨迹可能不同，直接比较电压并不能独立归因于量化噪声。"
                   "因此另取同一保守组真实轨迹的量化残差 q[k]=x[k]−y_ADC[k] 重放，"
                   "设置 Kp=Ki=0、Kd=0.06，只比较微分通道 α=0 与 α=0.85。"
                   "此残差是测量量化误差，区别于参考与真实位移之差的闭环跟踪误差。"),
        _table(["同轨迹量化残差重放（0.1～1.0 s）", "本次结果"], [
            ["未滤波微分电压 RMS / V", _number(probe_metrics.get("raw_derivative_rms_v"), 6)],
            ["α=0.85 微分电压 RMS / V", _number(probe_metrics.get("filtered_derivative_rms_v"), 6)],
            ["微分电压 RMS 降低 / %", _number(probe_metrics.get("reduction_percent"), 2)],
        ]),
        _picture("quantization_noise.png", "图 3　同一实际轨迹的 ADC 量化残差与微分通道重放，区分信号量化效应和闭环轨迹差异。"),
        _heading("6 习题 5-1：抗积分饱和对比"),
        _paragraph("两组输入均为 100 mm 阶跃，Kp=4.2、Ki=22.0、Kd=0.06，"
                   "输出限幅 ±10 V、Ts=1 ms、12 bit 传感器及相同物理模型。"
                   "主比较按教材标准差分取 α=0，不加积分分离；"
                   "补充比较让两组共同取 α=0.85，其他条件保持一致。"
                   "作业使用无工艺外载工况，以聚焦执行器饱和与内部状态结构。"),
        _paragraph(f"本次习题仿真时长为 {results['positional_raw']['config']['duration']:g} s，"
                   "长于课堂 S 型实验，以记录饱和、反向与后续响应；"
                   "末 200 ms 统计只描述尾段响应，不将其直接称为渐近稳态误差。"),
    ]
    blocks.extend(_homework_blocks(results, "raw", "6.1 教材式主比较", "anti_windup_raw.png"))
    blocks.extend(_homework_blocks(results, "filtered", "6.2 共同滤波补充比较", "anti_windup_filtered.png"))
    blocks += [
        _heading("6.3 状态记忆与退饱和解释", 2),
        _paragraph("位置式在正误差期间继续积累 I[k]。即使输出已经达到 +10 V，"
                   "其内部积分仍可能增长；当比例项或误差反向后，积累的 I[k] 会继续支撑正向输出，"
                   "形成退饱和迟滞和额外超调。限幅前候选电压和独立积分状态体现这一机制。"
                   "增量式把本周期截断后的实际输出写回 u[k−1]，不跨周期保留超出电压边界的候选值，"
                   "因此阻断这种无效输出累积；图中的积分增量并非位置式的全历史积分。"),
        _paragraph("对本次 100 mm 阶跃，初始误差为 10 cm。未滤波首步候选电压"
                   "为 4.2×10+22×10×0.001+0.06×10/0.001=642.22 V；"
                   "共同滤波 α=0.85 时为 42+0.22+0.15×600=132.22 V。"
                   "滤波或未滤波微分 kick 在下一步衰减，会产生明显负控制增量，"
                   "使增量式可能在误差仍为正时首次退出 +10 V，甚至转为负限幅。"
                   "所以课件中“一旦误差反向即可立即退出”必须理解为截断输出记忆减少 windup 的"
                   "结构特点，而不是本实验每次过零都会即时退饱和的保证；退出仍由总控制增量和状态决定。"),
        _paragraph("事件定义：误差过零取采样误差首次 ≤0 的采样点，并另列真实误差事件；"
                   "退出 +10 V 取相邻采样中从正饱和变为非正饱和的首个点。"
                   "另外列出过零之后发生的退出事件，避免将早期微分 kick 引起的退出误当作反号退饱和。"
                   "表中“—”表示仿真时窗内未观测到相应事件或不适用；所有事件时刻分辨率为 1 ms。"
                   "饱和总时长统计 ±10 V 的 ZOH 区间，末端单独采样不增加持续时间。"),
        _heading("7 指标验收与结论"),
    ]
    acceptance = []
    for name, label in zip(scurve_names, ["保守 + 滤波", "激进 + 滤波", "保守 + 无滤波"]):
        met = results[name]["metrics"]
        at1 = met.get("error_at_1s_mm")
        tail = met.get("final_200ms_max_abs_error_mm")
        acceptance.append([label, _number(at1),
                           "达到" if at1 is not None and abs(at1) < .05 else "未达到",
                           _number(tail),
                           "达到" if tail is not None and tail < .05 else "未达到"])
    blocks += [
        _table(["课堂组", "1.0 s 误差 / mm", "1.0 s |e|<0.05 mm", "末 200 ms 最大 |e| / mm", "保持窗 <0.05 mm"], acceptance),
        _paragraph("S 型参考在 1.0 s 到达 100 mm，不等于执行机构同一时刻完成精密定位。"
                   "上表将 1.0 s 瞬时真实误差与保持窗口最大真实误差分别验收。"
                   "12 bit 量化反馈的一个代码对应约 0.07324 mm，"
                   "某个量化读数不变化并不能证明真实位移稳定在 0.05 mm 以内。"
                   "有限 1.4 s 仿真也不等同于渐近稳定性证明。"),
        _paragraph("本实现完成了数字 PID、先滤波后增量、实际输出限幅记忆、七段解析轨迹、"
                   "1 kHz ZOH 闭环和非线性 Twin 联接，并按实际数据报告负载增益变化、"
                   "量化噪声抑制及习题 5-1 退饱和过程。"
                   "若固定 PID 未达到课件精密定位指标，应保留失败指标作为性能边界证据，"
                   "进一步研究第六章压力/流量非线性补偿与前馈或后续约束预测控制，"
                   "不能通过预设响应曲线替代物理仿真结果。"),
        _heading("8 复现与数据说明"),
        _paragraph("进入解压后的 chapter5 目录，执行 python -m pip install -r requirements.txt "
                   "安装依赖，再执行 python digital_pid_lab.py 重算全部七组实验。"
                   "默认 outputs/ 中生成 CSV/JSON、五幅图和本报告。"
                   "执行 python -m unittest -v test_chapter5 可运行实现检查。"
                   "控制采样间隔固定为 1 ms；CSV 的 time_s 包含 t=0 和终点。"
                   "SI 原始数据保存位置、速度、参考、腔压、量化测量、真实/采样误差、实际/候选电压、"
                   "微分项、积分状态或单步积分增量、阀芯状态、外载和 Kq。"
                   "summary.json 是报告定量结果的依据；quantization_probe 对同轨迹量化残差重放。"),
        _paragraph("材料依据：第五章课件第 3～8 页（离散 PID、微分滤波和硬限幅）、"
                   "第 9～14 页（S 型轨迹、离散时序和流量增益）、第 16 页（习题 5-1）；"
                   "第三章的阀口压差符号和阀控缸模型；第四章的参数辨识与时间基准对齐"
                   "（未提供测量数据，本次沿用标称参数）；"
                   "Word 的 DigitalPID 与 generate_scurve 两段代码。"
                   "原始 Twin 保留；闭环接口和补充假设在模型说明中明确列出。"),
    ]
    return blocks


def _write_markdown(blocks, output_dir):
    lines = []
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            lines.append("#" * (block["level"] + 1) + " " + block["text"])
        elif kind == "paragraph":
            lines.append(block["text"])
        elif kind == "table":
            escape = lambda s: str(s).replace("|", "\\|").replace("\n", "<br>")
            lines.append("| " + " | ".join(map(escape, block["headers"])) + " |")
            lines.append("| " + " | ".join("---" for _ in block["headers"]) + " |")
            lines.extend("| " + " | ".join(map(escape, row)) + " |" for row in block["rows"])
        elif kind == "picture" and (output_dir / block["filename"]).exists():
            lines.append(f"![{block['caption']}]({block['filename']})")
        lines.append("")
    path = output_dir / "实验报告.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def _set_font(run, size=None, bold=None):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = "Times New Roman"
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.append(fonts)
    fonts.set(qn("w:eastAsia"), "宋体")


def _three_line_table(table):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    table.autofit = True
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if edge in ("top", "bottom") else "nil")
        if edge in ("top", "bottom"):
            element.set(qn("w:sz"), "10")
        borders.append(element)
    for cell in table.rows[0].cells:
        props = cell._tc.get_or_add_tcPr()
        cb = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "5")
        cb.append(bottom)
        props.append(cb)
    # 标题行重复，避免长表跨页失去列名。
    row_props = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    row_props.append(repeat)
    for row in table.rows:
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))


def _write_docx(blocks, output_dir):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
    document = Document()
    document.core_properties.title = "第五章 数字 PID 闭环控制实验报告"
    document.core_properties.subject = "S 型跟踪、量化噪声、变负载机理及习题 5-1 抗积分饱和"
    document.core_properties.author = ""
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    for block in blocks:
        kind = block["type"]
        if kind == "heading":
            level = block["level"]
            p = document.add_paragraph() if level == 0 else document.add_heading(level=level)
            p.paragraph_format.keep_with_next = True
            if level == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(block["text"])
            _set_font(run, 17 if level == 0 else (14 if level == 1 else 12), True)
        elif kind == "paragraph":
            p = document.add_paragraph()
            _set_font(p.add_run(block["text"]), 10.5)
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block["headers"]))
            for cell, header in zip(table.rows[0].cells, block["headers"]):
                _set_font(cell.paragraphs[0].add_run(str(header)), 9, True)
            for values in block["rows"]:
                for cell, value in zip(table.add_row().cells, values):
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    _set_font(p.add_run(str(value)), 9)
            _three_line_table(table)
            document.add_paragraph().paragraph_format.space_after = Pt(3)
        elif kind == "picture":
            path = output_dir / block["filename"]
            if path.exists():
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(17))
                caption = document.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _set_font(caption.add_run(block["caption"]), 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    path = output_dir / "实验报告.docx"
    document.save(path)
    return path


def write_reports(results, probe_metrics, output_dir):
    """以相同内容生成两种报告，返回 {'markdown': path, 'docx': path}。

    results 包含三课堂组及四作业组，每组有 data/config/metrics。
    图像由主程序先生成；仅在图像已存在时插入，避免生成断链附件。
    本函数不启动仿真、不修改模型、也不从示意图构造实验数值。
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    blocks = _build_report(results, probe_metrics)
    markdown = _write_markdown(blocks, output_dir)
    docx = _write_docx(blocks, output_dir)
    return {"markdown": str(markdown), "docx": str(docx)}
