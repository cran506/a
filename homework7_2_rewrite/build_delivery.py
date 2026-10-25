#!/usr/bin/env python3
"""将实验报告、源代码和可核查原始记录打包；DOCX 转换需要 pandoc。"""
from pathlib import Path
import hashlib
import json
import shutil
import subprocess
from zipfile import ZipFile, ZIP_DEFLATED


ROOT=Path(__file__).resolve().parent


def main():
    validation=json.loads((ROOT/'outputs/validation.json').read_text(encoding='utf-8'))
    if not validation['passed']:
        raise RuntimeError('独立验证未通过，不生成正式交付包')
    summary=json.loads((ROOT/'outputs/summary.json').read_text(encoding='utf-8'))
    body=(ROOT/'REPORT.md').read_text(encoding='utf-8')
    start,end='<!-- AUTO_RESULTS_START -->','<!-- AUTO_RESULTS_END -->'
    rows=['| N | 视野 (ms) | 超调 (%) | P2 峰值 (MPa) | 最小两腔压力 (MPa) | 底层 QP 均值 (ms) | 底层 QP P95 (ms) |',
          '| ---: | ---: | ---: | ---: | ---: | ---: | ---: |']
    for n,m in summary['results'].items():
        rows.append(f'| {n} | {m["prediction_ms"]:.0f} | {m["overshoot_pct"]:.3f} | {m["peak_p2_mpa"]:.4f} | '
                    f'{m["min_pressure_mpa"]:.4f} | {m["backend_run"]["mean_ms"]:.3f} | {m["backend_run"]["p95_ms"]:.3f} |')
    rows+=['','| N | CVXPY 均值 (ms) | CVXPY P95 (ms) | 完整控制均值 (ms) | 完整控制 P95 (ms) | 初次 CVXPY 编译/求解 (ms) |',
           '| ---: | ---: | ---: | ---: | ---: | ---: |']
    for n,m in summary['results'].items():
        rows.append(f'| {n} | {m["cvxpy"]["mean_ms"]:.3f} | {m["cvxpy"]["p95_ms"]:.3f} | '
                    f'{m["controller"]["mean_ms"]:.3f} | {m["controller"]["p95_ms"]:.3f} | {m["warmup_cvxpy_ms"]:.3f} |')
    if start not in body or end not in body:
        raise RuntimeError('REPORT.md 缺少自动结果区标记')
    body=body[:body.index(start)+len(start)]+'\n\n'+'\n'.join(rows)+'\n\n'+body[body.index(end):]
    (ROOT/'REPORT.md').write_text(body,encoding='utf-8')
    pandoc=shutil.which('pandoc')
    if not pandoc:
        raise RuntimeError('生成 DOCX 需要安装 pandoc；REPORT.md 与实验仍可独立阅读运行')
    destination=ROOT/'deliverables';destination.mkdir(exist_ok=True)
    report=destination/'第七章_习题7-2_重写版报告.docx'
    subprocess.run([pandoc,str(ROOT/'REPORT.md'),'--from=markdown+tex_math_dollars',
                    '--to=docx',f'--resource-path={ROOT}','-o',str(report)],check=True,cwd=ROOT)
    from docx import Document
    from docx.shared import Inches,Pt
    from docx.oxml.ns import qn
    doc=Document(report)
    for section in doc.sections:
        section.page_width,section.page_height=Inches(11.69),Inches(8.27)
        section.top_margin=section.bottom_margin=Inches(.6)
        section.left_margin=section.right_margin=Inches(.65)
    for name in ('Normal','Body Text','Heading 1','Heading 2','Title'):
        if name not in doc.styles:continue
        style=doc.styles[name];style.font.name='Noto Sans CJK SC'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
        if name in ('Normal','Body Text'):style.font.size=Pt(10)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:run.font.size=Pt(9)
    for shape in doc.inline_shapes:
        ratio=shape.height/shape.width;shape.width=Inches(9.8);shape.height=int(shape.width*ratio)
    doc.save(report)
    archive=destination/'第七章_习题7-2_重写版完整作业.zip'
    files=[p for p in sorted(ROOT.rglob('*')) if p.is_file()
           and '__pycache__' not in p.parts and p.suffix not in ('.pyc','.zip','.sha256')
           and p.name!='manifest.json']
    hashes={p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    manifest=dict(folder=ROOT.name,validated=True,sha256=hashes)
    manifest_file=ROOT/'manifest.json'
    manifest_file.write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    with ZipFile(archive,'w',ZIP_DEFLATED) as z:
        for file in files+[manifest_file]:
            z.write(file,arcname=f'{ROOT.name}/{file.relative_to(ROOT).as_posix()}')
    with ZipFile(archive) as z:
        if z.testzip() is not None:raise RuntimeError('ZIP integrity check failed')
        for name,expected in hashes.items():
            if hashlib.sha256(z.read(f'{ROOT.name}/{name}')).hexdigest()!=expected:
                raise RuntimeError('ZIP checksum mismatch: '+name)
    digest=hashlib.sha256(archive.read_bytes()).hexdigest()
    archive.with_suffix('.zip.sha256').write_text(digest+'  '+archive.name+'\n',encoding='utf-8')
    print(report);print(archive);print('Verified files:',len(hashes))


if __name__=='__main__':
    main()
