"""用已有本次积分数据生成 Word 报告、SHA256 清单和自包含 ZIP。"""
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent
REPORT = '作业5-1_PID抗积分饱和实验报告.docx'
ARCHIVE = '作业5-1_PID抗积分饱和完整交付.zip'


def _number(value, places=3):
    return '未发生' if value is None else f'{value:.{places}f}'


def report():
    summary = json.loads((ROOT/'outputs/summary.json').read_text(encoding='utf-8'))
    validation = json.loads((ROOT/'outputs/validation.json').read_text(encoding='utf-8'))
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.65)
    style = document.styles['Normal']
    style.font.name = 'Noto Sans CJK SC'
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    document.add_heading('《智能流体传动与控制》作业 5-1', 0)
    document.add_paragraph('位置式与增量式 PID 抗积分饱和特性对比')
    document.add_heading('1. 实验条件与公式', 1)
    c = summary['positional']['config']
    p = c['physical_parameters']
    document.add_paragraph(
        f'两组使用相同液压数字孪生、初始状态、ADC、增益和执行器限幅。'
        f'阶跃指令在 t={c["step_time"]:.3f} s 从 0 跳到 {c["step_amplitude_m"]*1000:.0f} mm，'
        f'Ts={c["ts"]*1000:.0f} ms，仿真 {c["duration"]:.1f} s，输出限制 ±{c["voltage_limit_v"]:.0f} V。'
        f'PPT 给定 Kp={c["kp"]} V/cm、Ki={c["ki"]} V/(cm·s)、Kd={c["kd"]} V·s/cm；'
        f'α={c["alpha"]}。误差在控制器中以 cm 表示，图中以 mm 表示。')
    document.add_paragraph('实验 A：I[k]=I[k−1]+Ki·e[k]·Ts；'
                           'u_calc[k]=Kp·e[k]+I[k]+D[k]；u[k]=clip(u_calc[k],−10,+10)。')
    document.add_paragraph('实验 B：Δu[k]=Kp·(e[k]−e[k−1])+Ki·e[k]·Ts+D[k]−D[k−1]；'
                           'u_candidate[k]=u_actual[k−1]+Δu[k]；'
                           'u[k]=clip(u_candidate[k],−10,+10)，下一周期记忆写入此实际 u[k]。')
    document.add_paragraph(
        '默认两组共同采用测量值微分：D[k]=α·D[k−1]−(1−α)·Kd·(y[k]−y[k−1])/Ts。'
        '这是消除参考阶跃微分冲击的工程补充，阶跃瞬间并非原教材误差微分式的原样复现。'
        '原式对阶跃误差微分，首步 D=90 V、候选电压 132.22 V；保留 '
        'python work_5_1.py --derivative-mode error --no-plot 的独立对照数据。'
        '两组均没有积分分离、条件积分或额外电压平滑。')
    document.add_paragraph(
        f'阀控缸采用原有有向孔口流量、泄漏、压力连续性和一阶阀芯动态；'
        f'xv_max={p["xv_max"]*1000:.3f} mm（本作业公开的共同课堂整定；原示例 5.000 mm，'
        f'上一章轨迹实验 2.500 mm），阀时间常数 {c["valve_tau_s"]*1000:.1f} ms。'
        f'真实静摩擦 Fs={p["Fs"]:.0f} N、动摩擦 Fc={p["Fc"]:.0f} N，'
        f'零速时合力在 ±Fs 内保持静止，越过阈值后由事件求根释放。'
        f'ADC {c["adc_bits"]} bit，LSB={c["adc_lsb_mm"]:.6f} mm。无工艺外载。'
        '压力、位移、电压曲线均来自真实闭环积分，没有拼接参考曲线。')
    document.add_heading('2. 本次计算结果', 1)
    table = document.add_table(rows=1, cols=3)
    table.style = 'Light Shading Accent 1'
    for cell, value in zip(table.rows[0].cells, ('指标', '位置式 A', '增量式 B')):
        cell.text = value
    fields = [('overshoot_percent', '超调 / %', 3),
              ('peak_position_mm', '最大位移 / mm', 3),
              ('max_internal_voltage_v', '限幅前候选峰值 / V', 3),
              ('crossing_time_interpolated_s', '真实误差过零（插值）/ s', 6),
              ('first_nonpositive_sampled_error_s', '控制误差首个非正采样 / s', 3),
              ('initial_positive_saturation_exit_s', '首次退出 +10 V / s', 3),
              ('desaturation_delay_after_crossing_ms', '真实过零后退饱和迟滞 / ms', 3),
              ('final_error_mm', '1.2 s 真实误差 / mm', 4)]
    for key, label, places in fields:
        row = table.add_row().cells
        row[0].text = label
        for cell, kind in zip(row[1:], ('positional', 'incremental')):
            cell.text = _number(summary[kind]['metrics'][key], places)
    document.add_paragraph('过零插值时间、真实状态采样过零时间、ADC 控制误差过零时间分别记录。'
                           '增量式在误差反向前已退出正饱和，过零后的退饱和迟滞为 0；'
                           '不将提前退出的负时间差称为负迟滞。')
    document.add_picture(str(ROOT/'outputs/anti_windup_comparison.png'), width=Inches(6.5))
    document.add_heading('3. 积分饱和机理与参考图比较', 1)
    document.add_paragraph(
        '位置式只限制输出电压，显式积分 I 仍在饱和期间继续累加。误差变负时，'
        '积分中存留的正控制量仍可能大于负比例/微分项，造成退饱和迟滞及大幅超调。'
        '增量式将裁剪后的实际电压作为迭代记忆，丢弃超出执行器能力的增量；'
        '因此可在误差尚为正但 Δu<0 时提前减小输出。')
    document.add_paragraph(
        '定义增量组 I_eff[k]=u[k]−Kp·e[k]−D[k]，可得 '
        'I_eff[k]=I_eff[k−1]+Ki·e[k]·Ts+u[k]−u_candidate[k]。'
        '最后一项是限幅引入的跟踪修正。受到 ±10 V 约束的是实际输出记忆；'
        '候选控制量及 I_eff 不必位于此范围。若保存未限幅候选值而只裁剪送入执行器的电压，'
        '增量式同样可能 windup，因而不能将抗饱和归因于增量形式本身。')
    for kind, label in (('positional', '位置式'), ('incremental', '增量式')):
        differences = summary['reference_comparison'][kind]
        lines = [f'{key}：参考 {item["reference"]}，本次 {_number(item["actual"], 6)}'
                 for key, item in differences.items()]
        document.add_paragraph(label+'参考图对照：'+'；'.join(lines)+'。')
    document.add_paragraph('参考图的 31.4%、1.0%、42.5 V 和各时刻只作为比较目标，不作为控制器常数。'
                           '图形趋势可复现，精确数值取决于对象阀容量、腔容积、静摩擦与采样量化。'
                           '本次公开配置及两种微分模式数据均附在完整交付中。')
    document.add_heading('4. 验证与复现', 1)
    document.add_paragraph(
        f'CSV 独立检查及严格积分复算：{"通过" if validation["passed"] else "未全部通过"}。'
        f'9 项回归测试通过；检查数量 {len(validation["checks"])}，'
        f'基线 rtol={c["rtol"]:g}，严格复算 rtol=1e−7。')
    for kind, values in validation['numerical_comparison'].items():
        document.add_paragraph(kind+f' 最大位移差 {values["max_position_difference_mm"]:.6f} mm。')
    document.add_paragraph('运行：python -m pip install -r requirements.txt；python work_5_1.py；'
                           'python -m unittest -v；python verify_results.py --strict；'
                           'python package_delivery.py。具体文件与单位见 README.md。')
    destination = ROOT/'deliverables'/REPORT
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def included_files():
    return sorted(path for path in ROOT.rglob('*') if path.is_file()
                  and '__pycache__' not in path.parts and '.cache' not in path.parts
                  and path.suffix not in ('.zip', '.pyc') and path.name != 'manifest.json')


def package():
    report()
    files = included_files()
    manifest = {'project': 'homework5_1', 'files': {
        str(path.relative_to(ROOT)): {'bytes': path.stat().st_size,
                                    'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
        for path in files}}
    (ROOT/'manifest.json').write_text(json.dumps(manifest, indent=2, ensure_ascii=False),
                                     encoding='utf-8')
    archive = ROOT/'deliverables'/ARCHIVE
    with ZipFile(archive, 'w', ZIP_DEFLATED, compresslevel=9) as bundle:
        for path in [*files, ROOT/'manifest.json']:
            bundle.write(path, 'homework5_1/'+str(path.relative_to(ROOT)))
    with ZipFile(archive) as bundle:
        if bundle.testzip() is not None:
            raise RuntimeError('ZIP 完整性检查失败')
    print(archive)
    print(f'{len(files)+1} files; {archive.stat().st_size} bytes')


if __name__ == '__main__':
    package()
