"""从真实训练及闭环 CSV 生成第八章四联图和 Word 报告。

本模块没有参考曲线拟合或缺数据的模拟回退。缺少运行数据时明确报错。
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
from pathlib import Path
import tempfile

os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "chapter8_mpl"))
os.environ.setdefault("XDG_CACHE_HOME", str(Path(tempfile.gettempdir()) / "chapter8_cache"))
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np


POLICIES = ("pid", "nominal_ff", "sac_fp32", "edge_int8")
LABELS = {"pid": "传统 PID", "nominal_ff": "标称模型前馈", "sac_fp32": "SAC · FP32", "edge_int8": "70% 剪枝 + INT8"}
COLORS = {"pid": "#444444", "nominal_ff": "#ed8b00", "sac_fp32": "#3567ee", "edge_int8": "#15963e"}
STYLES = {"pid": "--", "nominal_ff": ":", "sac_fp32": "-", "edge_int8": "-."}


def _read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    if not rows:
        raise ValueError(f"没有真实数据：{path}")
    return rows


def _column(rows, name):
    try:
        return np.array([float(row[name]) for row in rows], dtype=float)
    except (KeyError, ValueError) as error:
        raise ValueError(f"CSV 字段 {name} 缺失或无效") from error


def _finite(x, name):
    if not np.isfinite(x).all():
        raise ValueError(f"{name} 含非有限数据")
    return x


def _number(value, digits=3):
    if value is None:
        return "未记录"
    if isinstance(value, (float, int)):
        return f"{value:.{digits}f}"
    return str(value)


def _recovery_time(t, e, fault_time=.5, band=1., hold=.05):
    post = np.flatnonzero(t >= fault_time - 1e-10)
    if not post.size:
        return None
    peak = post[np.argmax(np.abs(e[post]))]
    for i in range(peak, len(t)):
        end = np.searchsorted(t, t[i] + hold - 1e-10)
        if end < len(t) and np.all(np.abs(e[i:]) <= band):
            return float(t[i] - fault_time)
    return None


def _policy_metrics(rows):
    answer = {}
    for policy in POLICIES:
        group = [row for row in rows if row["policy"] == policy]
        if not group:
            raise ValueError(f"缺少实际闭环策略：{policy}")
        t = _finite(_column(group, "time_s"), policy + " time")
        e = _finite(_column(group, "velocity_error_mms"), policy + " error")
        u = _finite(_column(group, "control_v"), policy + " voltage")
        if np.any(np.diff(t) <= 0):
            raise ValueError(f"{policy} 时刻必须严格递增")
        post = t >= .5 - 1e-10
        tail = t >= t[-1] - .1 - 1e-10
        if not post.any():
            raise ValueError("闭环 CSV 没有 t≥0.5s 故障区")
        answer[policy] = {
            "minimum_velocity_error_mms": float(e[post].min()),
            "maximum_abs_velocity_error_mms": float(np.abs(e[post]).max()),
            "post_fault_rms_mms": float(np.sqrt(np.mean(e[post] ** 2))),
            "tail_mean_error_mms": float(np.mean(e[tail])),
            "recovery_s_band_1mms_hold_50ms": _recovery_time(t, e),
            "maximum_abs_voltage_v": float(np.abs(u).max()),
            "final_time_s": float(t[-1]),
            "sample_count": len(t),
        }
    return answer


def _set_font():
    available = {item.name for item in font_manager.fontManager.ttflist}
    candidates = ["Noto Sans CJK SC", "Noto Sans CJK JP", "WenQuanYi Micro Hei", "Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams.update({"font.family": "sans-serif", "font.sans-serif": [name for name in candidates if name in available],
                         "axes.unicode_minus": False, "font.size": 10, "axes.titlesize": 12,
                         "axes.labelsize": 10, "legend.fontsize": 8.5, "savefig.dpi": 180})


def _plot(out, training, evaluation, safety, safe_summary, measured):
    _set_font()
    fig, axes = plt.subplots(2, 2, figsize=(16.8, 9.7))
    fig.subplots_adjust(left=.065, right=.93, top=.88, bottom=.13, wspace=.28, hspace=.37)
    fig.suptitle("第八章 深度强化学习 SAC：真实训练、HOCBF 安全滤波与故障鲁棒性验证", fontsize=16, y=.967)
    for ax in axes.flat:
        ax.grid(True, linestyle=":", alpha=.4)
        ax.set_axisbelow(True)

    ax = axes[0, 0]
    steps = _finite(_column(training, "timesteps"), "training timesteps")
    rewards = _finite(_column(training, "reward_mean"), "training reward")
    entropy = _finite(_column(training, "voltage_entropy"), "training entropy")
    if np.any(np.diff(steps) <= 0):
        raise ValueError("训练日志步数必须严格递增")
    max_steps = steps[-1]
    ax.axvspan(0, min(50., max_steps / 1000), color="#eeeeee", zorder=0)
    if max_steps > 50000:
        ax.axvspan(50., min(400., max_steps / 1000), color="#e8edfc", zorder=0)
    if max_steps > 400000:
        ax.axvspan(400., max_steps / 1000, color="#e7f5e9", zorder=0)
    reward_line, = ax.plot(steps / 1000, rewards, color="#3567ee", linewidth=1.6, label="真实回合平均每步奖励")
    if np.max(np.abs(rewards)) > 10000:
        ax.set_yscale("symlog", linthresh=10)
        ax.set_ylabel("平均每步奖励（对称对数刻度）", color="#3567ee")
    else:
        ax.set_ylabel("平均每步奖励", color="#3567ee")
    ax.tick_params(axis="y", colors="#3567ee")
    right = ax.twinx()
    entropy_line, = right.plot(steps / 1000, entropy, color="#333333", linestyle="--", linewidth=1.5, label="电压动作微分熵 H(u)")
    right.set_ylabel("H(u) / nat（tanh 后，u 以 V 计）")
    ax.set_xlim(0, max_steps / 1000)
    ax.set_xlabel("实际训练交互步数 / 千步")
    ax.set_title(f"(a) SAC 训练记录（{int(max_steps):,} 步）", pad=12)
    ax.legend(handles=[reward_line, entropy_line], loc="best", framealpha=.95)
    ax.text(.02, .03, "原始日志直接绘制；阶段底色仅标记课件区间", transform=ax.transAxes, fontsize=8, color="#555555")
    late=steps>=400000
    if np.count_nonzero(late)>2:
        zoom=ax.inset_axes([.52,.13,.44,.29])
        zoom.plot(steps[late]/1000,rewards[late],color="#3567ee",lw=1.1)
        zoom.set_title("后期奖励：线性刻度放大",fontsize=7)
        zoom.tick_params(labelsize=6);zoom.grid(True,ls=":",alpha=.3)

    ax = axes[0, 1]
    t = _column(safety, "time_s")
    xs = _finite(_column(safety, "x_safe_mm"), "safe position")
    xu = _column(safety, "x_unsafe_mm")  # collision 后 NaN 是实验结束，不补画轨迹
    us = _finite(_column(safety, "u_safe_v"), "safe command")
    ax.axhspan(100, 115, color="#f9dce2", alpha=.8)
    ax.axhline(100, color="#df365c", linestyle="--", linewidth=1.2, label="100 mm 机械边界")
    unsafe_line, = ax.plot(t, xu, color="#e43c61", linestyle=":", linewidth=1.8, label="未滤波：首次触边终止")
    safe_line, = ax.plot(t, xs, color="#3567ee", linewidth=1.8, label="HOCBF-QP：受保护位移")
    ax.set_xlim(0, min(.25, float(t.max())))
    ax.set_ylim(15, 110)
    ax.set_xlabel("时间 t / s")
    ax.set_ylabel("活塞位移 / mm", color="#3567ee")
    right = ax.twinx()
    voltage_line, = right.plot(t, us, color="#208c38", linestyle="-.", linewidth=1.5, label="滤波后控制电压")
    right.set_ylabel("安全电压 / V", color="#208c38")
    right.tick_params(axis="y", colors="#208c38")
    right.set_ylim(min(-1., float(us.min()) - .6), max(10., float(us.max()) + .6))
    ax.set_title("(b) 二阶约化模型的采样 HOCBF 防撞实验", pad=12)
    ax.legend(handles=[unsafe_line, safe_line, voltage_line], loc="lower right", framealpha=.95)
    duration = float(safe_summary["duration_s"])
    text = (f"首次干预 {safe_summary['first_intervention_s']:.3f} s\n"
            f"最小电压 {safe_summary['minimum_safe_voltage_v']:.3f} V\n"
            f"{duration:g} s 位移 {safe_summary['safe_final_position_mm']:.6f} mm")
    ax.text(.55, .57, text, va="top", transform=ax.transAxes, fontsize=9,
            bbox=dict(boxstyle="round,pad=.35", facecolor="white", edgecolor="#3567ee", alpha=.95))

    for panel, key, ylabel, title in (
        (axes[1, 0], "velocity_error_mms", "速度跟踪误差 v − v_ref / (mm/s)", "(c) 5000 N 负载 + 体积模量减半：真实速度误差"),
        (axes[1, 1], "control_v", "实际阀控制电压 / V", "(d) FP32 与剪枝 INT8 的闭环电压对比"),
    ):
        panel.axvspan(.5, 1., color="#eeeeee", zorder=0)
        if key == "velocity_error_mms":
            panel.axhline(0, color="#555555", linewidth=.8, linestyle=":")
        else:
            panel.axhline(10, color="#b6b6b6", linewidth=.8, linestyle=":")
        for policy in POLICIES:
            group = [row for row in evaluation if row["policy"] == policy]
            panel.plot(_column(group, "time_s"), _column(group, key), color=COLORS[policy],
                       linestyle=STYLES[policy], linewidth=1.9 if policy == "sac_fp32" else 1.5, label=LABELS[policy])
        panel.set_xlim(.4, 1.)
        panel.set_xlabel("全局时间 t / s")
        panel.set_ylabel(ylabel)
        panel.set_title(title, pad=12)
        panel.legend(loc="best", framealpha=.95)
    ax = axes[1, 0]
    values = measured["sac_fp32"]
    rec = values["recovery_s_band_1mms_hold_50ms"]
    detail = (f"SAC 最大绝对误差 {values['maximum_abs_velocity_error_mms']:.2f} mm/s\n"
              + (f"恢复至 ±1 mm/s：{rec:.3f} s" if rec is not None else "本窗口未持续恢复至 ±1 mm/s"))
    ax.text(.46, .05, detail, transform=ax.transAxes, fontsize=8.5,
            bbox=dict(boxstyle="round,pad=.3", facecolor="white", edgecolor="#3567ee", alpha=.95))
    fig.text(.065, .067, "(a) 为实际 SAC 日志；(c)(d) 四策略共用四状态液压 ODE、初态及故障，确定性策略评估未叠加人工颤振。", fontsize=9, color="#444444")
    fig.text(.065, .041, f"(b) 采用独立约化模型 ẋ=v，v̇={safe_summary['k_gain_mm_s2_per_v']:g}u−{safe_summary['damping_per_s']:g}v；其证书不直接等同于四状态液压模型的安全证明。", fontsize=9, color="#444444")
    path = out / "chapter8_comparison.png"
    fig.savefig(path, facecolor="white")
    fig.savefig(path.with_suffix('.pdf'),facecolor='white')
    plt.close(fig)
    return path


def _flatten(obj, prefix="", max_items=90):
    result = []
    if isinstance(obj, dict):
        for key, value in obj.items():
            if key in {"history", "checks", "interval_minima"}:
                continue
            name = f"{prefix}.{key}" if prefix else str(key)
            if isinstance(value, (dict, list)):
                result.extend(_flatten(value, name, max_items))
            else:
                result.append((name, value))
            if len(result) >= max_items:
                break
    elif isinstance(obj, list) and len(obj) <= 12:
        result.append((prefix, ", ".join(map(str, obj))))
    return result[:max_items]


def _word_report(out, figure, training, safe_summary, measured, metadata, project):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml.ns import qn
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = section.bottom_margin = Inches(.7)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.add_heading("第八章 SAC、物理安全滤波与边缘量化实验报告", 0)
    document.add_paragraph("依据第八章课堂 PPT 与提供的三个模块完成。本文所有数值来自本目录的真实训练日志和仿真数据，不按参考图手工拼接曲线。")
    document.add_heading("实验范围与共同条件", level=1)
    settings = metadata.get("evaluation", {})
    start = float(settings.get("start_time_s", .4))
    display_start = float(settings.get("display_start_time_s", .4))
    fault = float(settings.get("fault_time_s", .5))
    document.add_paragraph(f"训练环境采用位置、速度、无杆腔压力、有杆腔压力四个液压状态。使用公开紧凑缸教学配置，使 5000N 对该对象构成显著负载；这套标度不是此前大面积液压缸的辨识参数。全部控制策略共用相同状态方程、初始巡航状态、采样周期与复合故障。评估从全局 {start:g}s 的共同平衡初态开始，图显示 {display_start:g}–1s，前段预运行同样保留在 CSV；{fault:g}s 施加负载并将油液体积模量乘以 0.5。")
    document.add_paragraph("传统组采用固定参数位置 PID，以巡航点电压为初始偏置，使用位置/速度误差反馈、不完全微分和条件积分抗饱和，具体增益及单位见 baseline_pid.py。标称前馈组仅使用空载平衡孔口/受力方程逆解，不含压力增益调度或积分修正，不能把其负载后静差写成‘前馈+PID’的表现。SAC 与 INT8 均读取相同观察量，INT8 独立闭环演化而不是重放 FP32 电压。")
    config = json.loads((project / "physical_config.json").read_text(encoding="utf-8"))
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "公开物理参数"
    table.rows[0].cells[1].text = "值"
    for key in ("A1", "A2", "L", "m", "Ps", "Pr", "beta_e", "F_c", "deadzone_v", "K_flow"):
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = str(config[key]) + "（SI 单位；死区为 V）"
    document.add_paragraph("对象具有有向孔口流量、可压缩两腔、内泄漏、低压补油与物理泄压。积分状态不作绘图裁剪。连续 tanh-Stribeck 摩擦近似不包含集合值静摩擦粘住分支，因此不能据此宣称验证了真实密封件的静摩擦爬行消除。")
    document.add_heading("实际 SAC 训练与熵的含义", level=1)
    final = training[-1]
    document.add_paragraph(f"实际交互步数为 {int(float(final['timesteps'])):,}。最后一条记录的回合平均每步奖励为 {float(final['reward_mean']):.6g}，电压动作微分熵为 {float(final['voltage_entropy']):.6g} nat。图中直接绘制记录值；课件三阶段底色不代表额外拟合出的收敛规律。")
    document.add_paragraph("归一化动作 a∈[−1,1]，控制电压 u=10a。训练熵从 tanh 压缩后的 actor 动作分布求取，换算 H(u)=H(a)+ln(10)，使用真实 replay buffer 状态上的蒙特卡洛估计。连续分布微分熵允许负值。随机预热期尚未由 actor 控制，此时曲线表示未更新 actor 的分布，不是均匀探索动作的行为熵。归一化动作最大熵为 ln(2)，所以不能把课件起始 2.8 nat 解释为该动作空间上的熵。熵目标若显式配置，属于设计选择，不能再描述成无条件自发出现的固定终值。")
    document.add_paragraph("训练奖励先将已结束回合的总奖励除以真实回合长度，再平均最近 30 个回合。若随机探索初期的原始奖励绝对值大于 10000，图 (a) 使用明确标注的对称对数刻度，以保留真实范围。训练随机策略的奖励统计与确定性故障测试分开报告。")
    document.add_paragraph("每步奖励按未归一化的毫米位置误差计算：r=−100e²−0.1u²−(u−u_prev)²+2·1{|e|<0.2mm}。训练优化器接收该奖励乘以 0.01 后的数值，以改善 critic 的数值标度；图表和 Monitor 日志仍保存原始奖励。熵温度是相对此缩放奖励调节的，因此这是明确的训练设计，不是逐项照搬未缩放目标。若环境触发明确失败终止，另加 −500 惩罚，条件以环境代码为准。奖励大小不能跨单位、跨回合长度直接与课件示意数值比较。")
    document.add_paragraph("状态扩展为十维：位置误差、速度误差、加速度误差、P1、P2、绝对位置、期望速度、期望加速度、上一实际电压、供油压力。压力和运动单位分别转换到 MPa、mm、mm/s 后按公开工程尺度标准化；这属于固定参考缩放，不能声称是从样本拟合出的严格零均值单位方差。随机参数本身仍可能隐藏，因而不宣称完全解决任意参数未知系统的部分可观测性。")
    normalization_path=out/'normalization_validation.json'
    if normalization_path.is_file():
        normalization=json.loads(normalization_path.read_text(encoding='utf-8'))
        document.add_paragraph(f"另外实现了数据拟合的经验 Z-score：用 {normalization['calibration_rows']} 条校准记录拟合均值和标准差，常量通道使用尺度 1；通过首层权重/偏置的仿射折叠保持原策略函数。在 {normalization['held_out_rows']} 条留出记录上，坐标变换前后最大归一化动作差为 {normalization['max_held_out_normalized_action_difference']:.3g}。这是训练后的等价坐标变换，原百万步训练历史仍使用原工程尺度。统计量及等价 actor 单独保存，未直接更换输入尺度而破坏已学控制律。")
    document.add_paragraph("域随机化覆盖质量 ±20%、初始油液体积模量 600–1400 MPa、摩擦 ±30%；若训练同时随机化负载/刚度故障，评估条件属于训练包络覆盖下的仿真鲁棒性验证。没有真实试验台数据，本项目没有完成真机迁移或真机小样本微调。")
    document.add_heading("四联结果图", level=1)
    document.add_picture(str(figure), width=Inches(9.8))
    document.add_heading("复合故障下的实际指标", level=1)
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    headings = ["策略", "最小误差 mm/s", "误差 RMS mm/s", "末段均值 mm/s", "恢复时间 s", "最大|u| V"]
    for cell, title in zip(table.rows[0].cells, headings):
        cell.text = title
    for policy in POLICIES:
        data = measured[policy]
        values = [LABELS[policy], _number(data['minimum_velocity_error_mms']), _number(data['post_fault_rms_mms']),
                  _number(data['tail_mean_error_mms']),
                  "未恢复" if data['recovery_s_band_1mms_hold_50ms'] is None else _number(data['recovery_s_band_1mms_hold_50ms']),
                  _number(data['maximum_abs_voltage_v'])]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    document.add_paragraph("速度误差定义为实际速度减参考速度；RMS 使用 t≥0.5s 的样本；末段均值使用本次记录最后 0.1s。恢复时间从故障时刻起计，在最大绝对误差之后寻找首次直至记录结束均满足 |误差|≤1mm/s、且留有至少 50ms 证据的时刻。未恢复如实列示。以上统计没有平滑原始响应。")
    fp = measured["sac_fp32"]
    edge_response = measured["edge_int8"]
    edge_recovery = edge_response['recovery_s_band_1mms_hold_50ms']
    fp_recovery=fp['recovery_s_band_1mms_hold_50ms']
    order=('早于' if fp_recovery is not None and edge_recovery is not None and edge_recovery<fp_recovery else '晚于或等于')
    edge_note = ("未满足直至窗口结束始终落入 ±1mm/s 的严格恢复判据" if edge_recovery is None
                 else f"按相同严格判据在 {edge_recovery:.3f}s 恢复，{order} FP32；该带宽判据不表示所有误差指标都更优")
    document.add_paragraph(f"与课件参考相比，本次 SAC 最大绝对速度误差为 {fp['maximum_abs_velocity_error_mms']:.3f}mm/s，大于示意图约 14.1mm/s；恢复时间 {_number(fp['recovery_s_band_1mms_hold_50ms'])}s 与参考 0.08s 为相同量级。控制电压先主动升高，再回到带载巡航值，复现了主要机制趋势，但峰值、回弹和时间历程没有逐点吻合。INT8 的故障后 RMS 为 {edge_response['post_fault_rms_mms']:.3f}mm/s，FP32 为 {fp['post_fault_rms_mms']:.3f}mm/s；INT8 仍有量化纹波，{edge_note}，不能称为两条曲线完全重合。")
    document.add_paragraph("FP32 与 INT8 采用确定性均值策略，因此输出不会凭空含有 SAC 的随机探索颤振。若需研究最大熵的概率穿越效应，必须另行采用带随机采样的策略并记录种子与动作频谱；本图不把人工正弦波叠加到网络输出。")
    document.add_heading("HOCBF 的适用模型及独立验证", level=1)
    document.add_paragraph(f"安全演示独立采用 ẋ=v、v̇={safe_summary['k_gain_mm_s2_per_v']:g}u−{safe_summary['damping_per_s']:g}v（mm、mm/s）的二阶课堂模型，机械边界为 0–100mm，内缩安全边界为 1.2–98.8mm。κ1={safe_summary['kappa1_per_s']:g}、κ2={safe_summary['kappa2_per_s']:g}，因此 α1={safe_summary['alpha_sum_per_s']:g}、α2={safe_summary['alpha_product_per_s2']:g}。参数由 safety_config.json 公开声明，不是四状态对象的辨识参数；传播模型、QP 与独立数值验证始终使用同一组系数。")
    document.add_paragraph("标量 QP 把策略电压投影到电压界与高阶屏障共同确定的可行区间；分析整个 1ms 零阶保持区间，并检查初始 h 和 ψ1 条件。不可行时明确报错，避免把随意保持输入误称安全。另用独立 DOP853 密输出复核随机可行状态及采样间不等式。")
    document.add_paragraph(f"本次首次干预为 {safe_summary['first_intervention_s']:.6f}s，最小安全电压为 {safe_summary['minimum_safe_voltage_v']:.6f}V，{safe_summary['duration_s']:g}s 末位移为 {safe_summary['safe_final_position_mm']:.9f}mm。未保护轨迹在 {safe_summary['unsafe_collision_time_s']:.9f}s 首次到达 100mm，之后记为 NaN（实验已终止），没有伪造碰撞后的平顶轨迹。")
    document.add_paragraph("这一二阶安全证书不能直接转移到训练所用四状态压力模型：位置对阀输入的相对阶数更高。图 (b) 用于验证给定约化模型上的最小干预原理，不能据此宣称真实液压装备获得了任意扰动下的绝对安全保证。")
    document.add_heading("剪枝、INT8 与推理时间", level=1)
    edge = metadata.get("edge", {})
    selection = edge.get("quantization_selection", {})
    selected_alpha = selection.get("selected_alpha", edge.get("validation", {}).get("smoothquant_alpha"))
    document.add_paragraph("对真实训练后 actor 的确定性均值网络按权重绝对值排序剪除约 70% 连接，再以同一 FP32 actor 为教师作固定掩码蒸馏，只调整保留连接。校准与验证使用不重叠的 replay 记录行，但来自同一训练分布，不能据此声称跨分布或真机独立验证。")
    document.add_paragraph(f"边缘导出采用 SmoothQuant 式输入/权重逐通道平衡：c 为校准输入均值，xmax_j=max|cal_j−c_j|，wmax_j=max|W_pruned[:,j]|，s_j=max(xmax_j,ε)^α/max(wmax_j,ε)^(1−α)，本次最终选用 α={_number(selected_alpha, 2)}、ε=1e−5。中心化输入幅度不超过 ε 的通道尺度设为 1。输入 z=(o−c)/s，首层折叠为 W′=W·diag(s)、b′=b+Wc；量化前与原稀疏网络代数等价，剪枝掩码及 FP32/pruned 路径不变。")
    document.add_paragraph("α 在最终完整校准集合上按 INT8 相对固定剪枝 FP32 actor 的动作 MAE，从 0.25/0.5/0.75 中选择。留出观察和故障测试的实际闭环输出不参与选参。此前缀不同于 Gym 固定工程尺度，中心和尺度随 v3 工件保存，推理不重新估计统计量。")
    document.add_paragraph("v3 使用逐输出通道对称 INT8 权重、仿射 INT8 隐层激活、INT32 点积与偏置累加、INT64 定点重缩放；最后线性层保留 INT32 logits，直到 tanh 边界才反量化。ReLU 层间正对角缩放在量化前代数等价，并保留所有稀疏零掩码。输入前缀、输入量化和末端 tanh 保留浮点，因此不称全路径纯整数，也不把最终 logits 误称 INT8。当前为 Python/NumPy CPU 实现；下表给出实际文件体积和计时，不套用课件 0.06ms 或 92.5% 压缩率。")
    document.add_paragraph("固定掩码压缩先执行 8000 步教师动作蒸馏，再执行 15000 步低学习率余弦 Adam 精修，总计 23000 步；只按校准损失保留结果，不增加 SAC 环境交互步数，也不依据闭环测试选择模型。")
    if edge:
        commission = edge.get("commissioning_calibration_count", 0)
        validation = edge.get("validation", {})
        document.add_paragraph(f"本次另外将 {commission} 个 FP32 单次 commissioning 轨迹观察加入激活范围校准。这条故障轨迹参与了校准，不是完全留出测试；INT8 部署仍使用自身反馈状态独立运行，没有重放 FP32 输出。下列动作误差在 {validation.get('evaluation_count', '未记录')} 条未参与校准的 replay 记录行上测量。")
        candidates = selection.get("candidates", [])
        if candidates:
            table = document.add_table(rows=1, cols=3)
            table.style = "Light Shading Accent 1"
            for cell, title in zip(table.rows[0].cells, ["SmoothQuant α 候选", "最终校准动作 MAE", "选择结果"]):
                cell.text = title
            for candidate in candidates:
                cells = table.add_row().cells
                cells[0].text = _number(candidate["alpha"], 2)
                cells[1].text = _number(candidate["calibration_action_mae"], 8)
                cells[2].text = "采用" if candidate["alpha"] == selected_alpha else "未采用"
            document.add_paragraph(f"候选评分使用 {selection.get('calibration_count', '未记录')} 条最终校准记录；完整选择过程保存为 quantization_calibration_selection.json。")
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Shading Accent 1"
        for cell, title in zip(table.rows[0].cells, ["工件 / CPU 参考路径", "实际文件字节", "中位耗时 ms", "P95 耗时 ms"]):
            cell.text = title
        for variant, label, filename in (
            ("fp32", "原始 FP32 actor", "actor_fp32.npz"),
            ("pruned_fp32", "剪枝蒸馏 FP32 actor", "actor_pruned_fp32.npz"),
            ("int8", "剪枝 INT8 actor", "actor_int8.npz"),
        ):
            benchmark = edge.get("benchmark", {}).get(variant, {})
            size = edge.get("artifact_bytes", {}).get(filename)
            values = [label, str(size) if size is not None else "未记录",
                      _number(benchmark.get("median_ms"), 6), _number(benchmark.get("p95_ms"), 6)]
            for cell, value in zip(table.add_row().cells, values):
                cell.text = value
        document.add_paragraph("文件字节为实际 .npz 序列化工件大小（包含元数据及压缩开销）；耗时为当前 CPU 的 batch=1 重复测量，不能作为 MCU 实测指标。非结构化零权重仍由此参考实现按稠密数组存放和计算，所以剪枝比例不直接等于速度或文件体积减少比例。")
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"
        table.rows[0].cells[0].text = "压缩质量指标"
        table.rows[0].cells[1].text = "本次实际值"
        sparsity = validation.get("actual_weight_sparsity")
        compression_values = [
            ("实际零权重比例", f"{100. * sparsity:.4f}%" if sparsity is not None else "未记录"),
            ("留出动作 MAE：剪枝 FP32 / 原始 FP32", _number(validation.get("pruned_action_mae"), 8)),
            ("留出动作 MAE：INT8 / 剪枝 FP32", _number(validation.get("int8_vs_pruned_action_mae"), 8)),
            ("留出动作 MAE：INT8 / 原始 FP32", _number(validation.get("int8_vs_fp32_action_mae"), 8)),
            ("留出动作最大绝对误差：INT8 / 原始 FP32", _number(validation.get("int8_vs_fp32_action_max_abs_error"), 8)),
            ("两次闭环速度 RMS 差（完整 0.3–1s）/ mm/s", _number(edge.get("closed_loop_velocity_rms_difference_mms"), 6)),
            ("两次闭环速度 RMS 差（显示 0.4–1s）/ mm/s", _number(edge.get("displayed_window_velocity_rms_difference_mms"), 6)),
            ("输入量化超范围比例", _number(validation.get("input_saturation_fraction"), 8)),
            ("输入平衡方法", str(validation.get("input_balancing_mode", "未记录"))),
            ("SmoothQuant α", _number(validation.get("smoothquant_alpha"), 4)),
            ("固定掩码蒸馏优化步数", str(edge.get("distillation", {}).get("steps", "未记录"))),
        ]
        for label, value in compression_values:
            cells = table.add_row().cells
            cells[0].text = label
            cells[1].text = value
        document.add_paragraph("动作误差以归一化动作 a 为单位，乘以 10 得到电压误差 V。闭环速度 RMS 差比较两条独立部署轨迹，不等同于同一观察上的网络动作误差；完整区间包含预运行，显示区间与图 (c)(d) 一致。")
    for key in ("training",):
        if key not in metadata:
            continue
        document.add_heading("运行元数据：" + key, level=2)
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Shading Accent 1"
        table.rows[0].cells[0].text = "字段"
        table.rows[0].cells[1].text = "实测值 / 配置"
        for field, value in _flatten(metadata[key], max_items=45):
            cells = table.add_row().cells
            cells[0].text = field
            cells[1].text = _number(value, 6)
    document.add_heading("复现与原始数据", level=1)
    validation_path = out / "physics_validation.json"
    if validation_path.exists():
        numerical = json.loads(validation_path.read_text(encoding="utf-8"))
        error = numerical.get("max_absolute_state_error_si_by_substeps", {}).get("8")
        if error is not None and len(error) == 4:
            document.add_paragraph(f"独立数值复核：在 {numerical['comparison_duration_s']:g}s 输入/故障重放中，native 8 子步 RK4 相对独立 Radau 的最大差为位移 {error[0]*1000:.6g}mm、速度 {error[1]*1000:.6g}mm/s、P1 {error[2]/1e6:.6g}MPa、P2 {error[3]/1e6:.6g}MPa；16/32 子步结果进一步收敛。该验证针对所记录的试验输入，详细检查与容差见 physics_validation.json。")
    document.add_paragraph("运行入口为 drl_hydraulic_lab.py；训练过程、评估、压缩和绘图分离。材料目录保留原 PPT。outputs/training.csv、evaluation.csv、safety.csv 和 JSON 摘要分别记录训练、故障评估、安全实验与元数据；models/ 保存真实 SAC 和 INT8 actor。完整命令及依赖见 README.md。")
    path = out / "第八章实验报告.docx"
    document.save(path)
    return path


def build_report(project_dir=None):
    project = Path(project_dir or Path(__file__).resolve().parent)
    out = project / "outputs"
    training = _read_csv(out / "training.csv")
    evaluation = _read_csv(out / "evaluation.csv")
    safety = _read_csv(out / "safety.csv")
    metadata = json.loads((out / "metrics.json").read_text(encoding="utf-8"))
    safe_summary = json.loads((out / "safety_summary.json").read_text(encoding="utf-8"))
    recorded_steps = int(float(training[-1]["timesteps"]))
    declared_steps = metadata.get("training", {}).get("actual_timesteps")
    if declared_steps is not None and recorded_steps != int(declared_steps):
        raise ValueError("训练日志末步与模型元数据不一致，拒绝混用不同运行的结果")
    common_times = None
    for policy in POLICIES:
        rows = [row for row in evaluation if row["policy"] == policy]
        if not rows:
            raise ValueError(f"缺少实际策略：{policy}")
        times = _column(rows, "time_s")
        if common_times is None:
            common_times = times
        elif times.shape != common_times.shape or not np.allclose(times, common_times, rtol=0., atol=1e-10):
            raise ValueError("四种策略的评估时间网格不一致")
        if np.max(np.abs(_finite(_column(rows, "control_v"), policy + " voltage"))) > 10. + 1e-6:
            raise ValueError(f"{policy} 记录电压超出已声明的执行器限值")
    if not safe_summary.get("passed", False):
        raise ValueError("HOCBF 独立验证未通过，不能生成通过验证的结果报告")
    measured = _policy_metrics(evaluation)
    figure = _plot(out, training, evaluation, safety, safe_summary, measured)
    report = _word_report(out, figure, training, safe_summary, measured, metadata, project)
    (out / "report_metrics.json").write_text(json.dumps(measured, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    return {"figure": str(figure), "report": str(report), "computed_policy_metrics": measured}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-dir", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    print(json.dumps(build_report(args.project_dir), ensure_ascii=False, indent=2))
