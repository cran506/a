"""从已有第七章结果生成 Word 报告、SHA256 清单及完整交付 ZIP。"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
REPORT = "第七章_MPC约束控制实验报告.docx"
ARCHIVE = "第七章_MPC约束控制完整交付.zip"
KINDS = ("pid", "input_mpc", "full_mpc")
NAMES = {"pid": "前馈 + PID", "input_mpc": "仅输入约束 MPC", "full_mpc": "输入 + P2 约束 MPC"}
EXCLUDE_DIRS = {"__pycache__", ".cache", ".pytest_cache", ".mypy_cache", ".ruff_cache", ".git"}


def _load(path):
    return json.loads(path.read_text(encoding="utf-8"))


def _number(value, places=3):
    if value is None:
        return "未记录"
    if isinstance(value, bool):
        return "通过" if value else "未通过"
    return f"{float(value):.{places}f}"


def _results(summary):
    values = summary.get("results", summary)
    missing = set(KINDS) - set(values)
    if missing:
        raise RuntimeError(f"summary.json 缺少策略：{sorted(missing)}")
    for kind in KINDS:
        if "metrics" not in values[kind]:
            raise RuntimeError(f"summary.json 缺少 {kind}.metrics")
    return values


def _figure(folder, requested=None):
    if requested is not None:
        path = Path(requested)
        path = path if path.is_absolute() else ROOT / path
    else:
        candidates = sorted(folder.glob("*.png"))
        if len(candidates) != 1:
            raise RuntimeError("outputs 必须有唯一结果 PNG，或用 --figure 指定已有四联图")
        path = candidates[0]
    if not path.is_file():
        raise RuntimeError(f"结果图尚未生成：{path}")
    if not path.resolve().is_relative_to(ROOT.resolve()):
        raise RuntimeError("结果图须放在本章目录内，确保交付包可以独立复现")
    return path


def _three_line_table(document, headers):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single" if edge in ("top", "bottom") else "nil")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "505050")
        borders.append(node)
    properties.append(borders)
    for cell in table.rows[0].cells:
        properties = cell._tc.get_or_add_tcPr()
        border = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        border.append(bottom)
        properties.append(border)
    return table


def _config_text(config, key, multiplier=1., default=None):
    value = config.get(key, default)
    return "未记录" if value is None else f"{float(value)*multiplier:g}"


def _validation_status(validation):
    for key in ("passed", "all_passed", "all_checks_passed"):
        if key in validation and isinstance(validation[key], bool):
            return "全部通过" if validation[key] else "未全部通过"
    return "以 validation.json 各项记录为准（未记录统一通过标志）"


def report(summary, validation, figure, folder):
    results = _results(summary)
    config = results["full_mpc"].get("config", summary.get("config", {}))
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.65)
        section.left_margin = section.right_margin = Inches(.65)
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = "Noto Sans CJK SC"
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    document.styles["Normal"].font.size = Pt(10)
    document.add_heading("第七章：模型预测约束控制与滚动优化", 0)
    document.add_paragraph("50 mm 极限定位：前馈 PID、输入约束 MPC 与输入及 P2 约束 MPC 三组课堂对比")
    document.add_heading("1. 实验条件与材料范围", 1)
    document.add_paragraph(
        f"三策略使用共同真实液压对象、初始状态和参考指令。指令在 "
        f"t={_config_text(config, 'step_time')} s 跃迁至 {_config_text(config, 'target_mm')} mm；"
        f"Ts={_config_text(config, 'ts', 1000.)} ms，"
        f"时长 {_config_text(config, 'duration')} s，预测时域 N={config.get('horizon', '未记录')}。"
        "参数是公开的课堂仿真假设，未由设备测量数据辨识。实际参数以本包配置及 summary.json 为准。")
    document.add_paragraph(
        "材料依据为 PPT 第 7.4 节三策略课堂实验。PPT 的超调和压力示意数值只用于趋势比较，"
        "未写入控制律或压力状态。课后习题 7-2 的 N=5/20/50 独立时域对比不在本次三策略实验范围内。"
        "原上传 Twin 保留，真实闭环扩展、预测模型及参数覆盖在独立文件中，文件 SHA256 见 manifest.json。")
    if config:
        keys = ("valve_tau", "makeup_area", "rtol", "max_du_step", "pressure_margin", "mpc_outer_iterations",
                "pid_kp", "pid_ki", "pid_kd", "pid_alpha", "ff_velocity", "ff_accel_time")
        recorded = [f"{key}={config[key]}" for key in keys if key in config]
        if recorded:
            document.add_paragraph("本次控制配置：" + "；".join(recorded) + "。")
    physical = config.get("physical_parameters", {})
    if physical:
        conditions = [("L", "行程", 1., "m"), ("V0_1", "无杆腔死容积", 1e6, "cm³"),
                      ("V0_2", "有杆腔死容积", 1e6, "cm³"),
                      ("beta_e", "标称体积模量", 1e-9, "GPa"),
                      ("xv_max", "最大阀开度", 1e3, "mm"), ("mass", "质量", 1., "kg"),
                      ("Fs", "静摩擦", 1., "N"), ("Fc", "动摩擦", 1., "N")]
        description = [f"{label} {float(physical[key])*scale:g} {unit}"
                       for key, label, scale, unit in conditions if key in physical]
        document.add_paragraph("共同对象参数：" + "；".join(description) + "。")
    if "initial_state" in config:
        document.add_paragraph("共同初态（mm、m/s、MPa、MPa、V）："
                               + json.dumps(config["initial_state"], ensure_ascii=False)
                               + "。预充压力按净液压力机械平衡初始化；内泄漏仍存在，不宣称零流量恒压平衡。")
    document.add_heading("2. 预测与执行机制", 1)
    document.add_paragraph(
        "物理状态为 [相对位置 mm、速度 m/s、P1 MPa、P2 MPa、阀位等效 V]，"
        "加入上一周期实际电压 u_prev 后为 6 维增广状态，优化变量为未来电压增量。"
        "局部雅可比与矩阵指数构建含仿射项的 LTV 零阶保持离散模型，保留初始状态及绝对压力偏置；"
        "有限时域二次规划以 OSQP 为主要求解器，只执行首个电压，下一采样以真实状态重新规划。")
    if all(key in config for key in ("q_position", "q_velocity", "move_weight", "settling_move_weight",
                                      "settling_error_mm", "settling_velocity_m_s")):
        document.add_paragraph(
            f"两组 MPC 共用 Q=diag([{config['q_position']:g},{config['q_velocity']:g},0,0,0])，"
            f"分别惩罚位置 mm 与速度 m/s 的误差；运动期电压增量权重 W={config['move_weight']:g}。"
            f"当 |位置误差|≤{config['settling_error_mm']:g} mm 且 "
            f"|速度|≤{config['settling_velocity_m_s']:g} m/s 时，"
            f"两组共同切换到 W={config['settling_move_weight']:g}，抑制末端指令抖动。"
            "这是基于当前状态的优化代价调度，未平滑计算后的电压或轨迹；"
            "CSV 的 move_weight、settling_cost_active 与逐步日志保存实际模式。")
    document.add_paragraph(
        "输入约束组限制实际电压 ±10 V 与单步电压变化量；完整组在相同输入约束上增加 P2 上限。"
        f"名义安全线为 {_config_text(config, 'p_safe', default=21.)} MPa，"
        f"预测裕量为 {_config_text(config, 'pressure_margin')} MPa。"
        "默认完整组没有开启 P1 上限或两腔压力下限，因此不能据 P2 结果宣称两腔均受硬保护。"
        "局部 QP 解还需检查残差及非线性预测可行性。")
    solver_settings = config.get("mpc_solver_settings", {})
    reserve = solver_settings.get("linearization_reserve_mpa")
    if reserve is not None and all(key in config for key in ("p_safe", "pressure_margin")):
        nonlinear_limit = config["p_safe"] - config["pressure_margin"]
        document.add_paragraph(
            f"完整组的非线性预测节点上界为 {nonlinear_limit:g} MPa；"
            f"局部 QP 再向内缩 {reserve:g} MPa，使用 {nonlinear_limit-reserve:g} MPa 上界，"
            "为局部线性化曲率保留数值余量。实际安全线仍为 21 MPa，"
            "不能将预测节点界与真实采样间压力峰混为同一个量。")
    document.add_paragraph(
        "OSQP 的状态、有限值及独立原始/对偶残差通过后才接受。若未通过，备用 Clarabel "
        "解同一个 QP 并独立验证；原 OSQP 失败状态保留，不重标为成功。"
        "历史移位方案只有从当前测量状态重新检查输入边界与非线性压力可行性后才能采用；"
        "没有可信方案即明确报错，不静默持位、软化压力限制或强制裁剪压力状态。")
    document.add_paragraph(
        "真实对象采用事件驱动静摩擦、硬阀死区、有向孔口、泄漏、补油及阀动态的混合 ODE；"
        "预测对象采用公开的光滑滑动近似，使雅可比和局部优化可计算。预测正则化只属于控制模型，"
        "不会修改真实摩擦事件或真实死区。图中的压力来自高密积分轨迹，包含控制采样间的压力峰；"
        "未人为注入压力脉冲、裁剪压力状态或平滑绘图。局部模型约束及有限场景验证不构成全局安全保证。")
    regularization = config.get("prediction_regularization", {})
    if regularization:
        document.add_paragraph("预测正则化尺度：" + "；".join(f"{key}={value}"
                               for key, value in regularization.items()) + "。")
    document.add_heading("3. 本次实际结果", 1)
    table = _three_line_table(document, ("指标", *[NAMES[k] for k in KINDS]))
    metrics = (("peak_position_mm", "最大相对位移 / mm", 4),
               ("overshoot_percent", "位置超调 / %", 3),
               ("first_positive_saturation_exit_s", "首次退出正饱和 / s", 3),
               ("first_target_crossing_s", "真实误差过零（采样插值）/ s", 6),
               ("final_error_mm", "0.15 s 真实误差 / mm", 4),
               ("peak_p2_mpa", "高密真实 P2 峰值 / MPa", 5),
               ("constraint_violation_p2_mpa", "超过 P2 安全线的峰值 / MPa", 5),
               ("peak_p1_mpa", "真实 P1 峰值（未约束）/ MPa", 5),
               ("max_internal_voltage_v", "限幅前候选电压峰值 / V", 4),
               ("max_abs_voltage_v", "最大实际 |u| / V", 4),
               ("max_abs_control_increment_v", "最大实际 |Δu| / V", 4),
               ("last_20ms_max_abs_error_mm", "最后 20 ms 最大 |e| / mm", 4),
               ("qp_runtime_median_ms", "QP 实测耗时中位数 / ms", 4),
               ("qp_runtime_max_ms", "QP 实测最大耗时 / ms", 4),
               ("total_controller_median_ms", "完整控制步骤耗时中位数 / ms", 4),
               ("total_controller_max_ms", "完整控制步骤最大耗时 / ms", 4),
               ("controller_within_1ms_fraction", "完整步骤 ≤1 ms 的比例 / %", 2),
               ("validated_shifted_plan_count", "经非线性复核的移位方案次数", 0),
               ("solved_inaccurate_count", "solved inaccurate 记录次数", 0))
    for key, label, places in metrics:
        cells = table.add_row().cells
        cells[0].text = label
        for cell, kind in zip(cells[1:], KINDS):
            value = results[kind]["metrics"].get(key)
            if key == "controller_within_1ms_fraction" and value is not None:
                value *= 100.
            cell.text = "不适用" if kind == "pid" and key.startswith("qp_") else _number(value, places)
    document.add_paragraph(
        "超调以本次真实位移峰值及目标位移计算；P2 越限量是实际压力相对安全线的正部，"
        "不以控制采样压力替代积分峰值。P1 峰值单独披露，用于说明保护范围。"
        "图例和表格均引用本次记录，未将 PPT 示意峰值当作计算结果。")
    baseline = results["pid"]["metrics"]
    exit_time, crossing_time = (baseline.get(key) for key in
                                ("first_positive_saturation_exit_s", "first_target_crossing_s"))
    if exit_time is not None and crossing_time is not None:
        ordering = "早于" if exit_time < crossing_time else "晚于或等于"
        document.add_paragraph(
            f"传统组首次退出正饱和在 {exit_time:.3f} s，{ordering}真实误差过零 "
            f"{crossing_time:.6f} s。"
            + ("因此本次不能将其超调归因于误差反向后的正饱和迟滞；"
               "高速惯性、阀动态、计划前馈及累积积分共同影响超调。"
               if exit_time < crossing_time else
               "事件次序按实际记录保留，退饱和机制不能只凭波形示意推断。"))
    document.add_paragraph(
        f"本次传统组在 {_config_text(config, 'duration')} s 的真实误差为 "
        f"{_number(baseline.get('final_error_mm'), 4)} mm。"
        "非对称活塞面积、内泄漏与闭阀存压使本真实模型尾段仍保持较高压力；"
        "未人为把压力恢复到示意图的 2 MPa。末端残差与尾段压力形状也属于公开差异。")
    reference = summary.get("reference_comparison", {})
    if reference:
        comparison = _three_line_table(document, ("策略", "PPT 超调 / %", "实际超调 / %",
                                                  "PPT P2 峰值 / MPa", "实际 P2 峰值 / MPa"))
        for kind in KINDS:
            item = reference.get(kind, {})
            actual = results[kind]["metrics"]
            cells = comparison.add_row().cells
            texts = (NAMES[kind], _number(item.get("reference_overshoot_percent")),
                     _number(actual.get("overshoot_percent")),
                     _number(item.get("reference_peak_p2_mpa")), _number(actual.get("peak_p2_mpa")))
            for cell, value in zip(cells, texts):
                cell.text = value
        document.add_paragraph("上述差异保留公开。对象参数、预测近似、摩擦及优化配置会改变实测值，"
                               "不保证示意图逐点重合，也不通过参数写死某个压力峰。")
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    document.add_picture(str(figure), width=width)
    document.add_paragraph("图 1　三策略实际位移、电压、有杆腔高密压力与速度；彩色点标记本次轨迹极值。")
    document.add_heading("4. 求解耗时与验证边界", 1)
    document.add_paragraph(
        "Ts=1 ms 表示仿真控制采样周期。QP 的计时与完整控制步骤计时分别记录，"
        "后者还包括局部线性化、预测、可行性检查等开销。中位数不能代替最坏时耗；"
        "未据本次桌面运行宣称满足 1 ms 硬实时。具体计时受 CPU、线程、软件版本与系统负载影响。")
    input_time = results["input_mpc"]["metrics"].get("total_controller_median_ms")
    full_time = results["full_mpc"]["metrics"].get("total_controller_median_ms")
    if input_time is not None and full_time is not None:
        budget = float(config.get("ts", .001)) * 1000.
        conclusion = (f"两组中位数均超过 {budget:g} ms 采样预算，当前实现未达到该实时要求。"
                      if min(input_time, full_time) > budget else
                      "是否满足实时要求须结合最大耗时、超预算比例和目标硬件进一步判断。")
        document.add_paragraph(
            f"本次仅输入组、完整组的完整控制步骤中位数分别为 "
            f"{input_time:.3f}、{full_time:.3f} ms。" + conclusion)
    document.add_paragraph("功能与数值检查状态：" + _validation_status(validation) + "。原始检查条件及独立复算见 outputs/validation.json。")
    document.add_paragraph(
        "严格积分复算使用本次已记录的控制电压，以更严格的积分容差重放真实对象，"
        "用于检查物理轨迹及采样间压力峰值的数值稳定性；该复算不重新执行 MPC 优化闭环。")
    for key in ("test_count", "tests_run", "unit_test_count"):
        if key in validation and isinstance(validation[key], (int, float)):
            document.add_paragraph(f"记录的回归测试数量：{int(validation[key])}。")
            break
    checks = validation.get("checks")
    if isinstance(checks, dict):
        document.add_paragraph(f"验证文件包含 {len(checks)} 项具名检查；不将未执行的检查标为通过。")
        failed = [name for name, value in checks.items() if value is False or
                  (isinstance(value, dict) and value.get("passed") is False)]
        if failed:
            document.add_paragraph("未通过项：" + "、".join(failed) + "。")
    document.add_heading("5. 复现与交付文件", 1)
    entry = "mpc_hydraulic_lab.py"
    if not (ROOT / entry).is_file():
        entries = sorted(ROOT.glob("*lab.py"))
        entry = entries[0].name if len(entries) == 1 else "README.md 中的实际实验入口"
    document.add_paragraph(
        "解压后进入 chapter7_latest 目录：python -m pip install -r requirements.txt；"
        f"python {entry}；python -m unittest -v；"
        "独立复核命令按 README.md 执行；python package_delivery.py 重新生成报告、清单和 ZIP。")
    document.add_paragraph(
        "交付包括本章全部源码、参数 JSON、依赖、教材材料、三组控制采样 CSV、高密压力数据、"
        "summary.json、validation.json、唯一四联图及本报告。CSV 单位直接编码在字段名中。"
        "manifest.json 按相对路径记录各文件字节数与 SHA256；ZIP 自身和 manifest 自身不参与清单哈希，"
        "防止循环依赖。相同设置可复现实验机制，计时及报告元数据不承诺逐字节一致。")
    destination = ROOT / "deliverables" / REPORT
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def included_files():
    return sorted((path for path in ROOT.rglob("*") if path.is_file()
                   and not EXCLUDE_DIRS.intersection(path.relative_to(ROOT).parts)
                   and path.name != "manifest.json"
                   and path.suffix.lower() not in {".zip", ".pyc", ".pyo", ".tmp"}),
                  key=lambda path: path.relative_to(ROOT).as_posix())


def package(figure=None):
    folder = ROOT / "outputs"
    required = {"README.md", "requirements.txt", "physical_config.json", "classroom_config.json",
                "experiment_config.json", "hydraulic_digital_twin.py", "hydraulic_plant.py", "mpc_core.py",
                "plotting.py", "outputs/summary.json", "outputs/validation.json",
                "outputs/pid.csv", "outputs/input_mpc.csv", "outputs/full_mpc.csv",
                "outputs/pid_pressure_trace.csv", "outputs/input_mpc_pressure_trace.csv",
                "outputs/full_mpc_pressure_trace.csv",
                "materials/chap07-ppt.pdf"}
    available = {path.relative_to(ROOT).as_posix() for path in included_files()}
    if not required <= available:
        raise RuntimeError(f"交付文件尚未齐全：{sorted(required - available)}")
    image = _figure(folder, figure)
    report(_load(folder / "summary.json"), _load(folder / "validation.json"), image, folder)
    files = included_files()
    manifest = {"project": "chapter7_latest", "files": {
        path.relative_to(ROOT).as_posix(): {
            "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        for path in files}}
    manifest_path = ROOT / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    destination = ROOT / "deliverables" / ARCHIVE
    with ZipFile(destination, "w", ZIP_DEFLATED, compresslevel=9) as bundle:
        for path in [*files, manifest_path]:
            bundle.write(path, (Path(ROOT.name) / path.relative_to(ROOT)).as_posix())
    with ZipFile(destination) as bundle:
        damaged = bundle.testzip()
        if damaged is not None:
            raise RuntimeError(f"ZIP 完整性检查失败：{damaged}")
        for relative, item in manifest["files"].items():
            payload = bundle.read(f"{ROOT.name}/{relative}")
            if len(payload) != item["bytes"] or hashlib.sha256(payload).hexdigest() != item["sha256"]:
                raise RuntimeError(f"ZIP 文件与清单不符：{relative}")
    print(f"报告：{ROOT / 'deliverables' / REPORT}")
    print(f"ZIP：{destination}（{len(files)+1} 个文件，{destination.stat().st_size} 字节）")
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--figure", help="指定章目录内已有四联结果 PNG；默认 outputs 中唯一 PNG")
    args = parser.parse_args()
    package(args.figure)


if __name__ == "__main__":
    main()
