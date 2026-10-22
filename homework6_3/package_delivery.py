"""用作业6-3已有结果生成 Word 报告、SHA256 清单和完整 ZIP，不执行仿真。"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
REPORT_NAME = "第六章作业6-3_死区补偿报告.docx"
ARCHIVE_NAME = "第六章作业6-3_完整交付.zip"
KINDS = ("under", "hard_over", "smooth")
LABELS = {"under": "A 欠补偿", "hard_over": "B 硬过补偿", "smooth": "C tanh 光滑补偿"}
EXCLUDED_DIRS = {"__pycache__", ".cache", ".pytest_cache", ".mypy_cache", ".ruff_cache", ".git"}


def _json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def _number(value, places=4):
    if value is None:
        return "未记录"
    if isinstance(value, bool):
        return "通过" if value else "未通过"
    number = float(value)
    if round(number, places) == 0:
        number = 0.
    return f"{number:.{places}f}"


def _table(document, headers):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = True
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if edge in ("top", "bottom") else "nil")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:color"), "555555")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for cell in table.rows[0].cells:
        properties = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        element = OxmlElement("w:bottom")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        borders.append(element)
        properties.append(borders)
    return table


def _status(validation):
    for key in ("passed", "all_passed", "all_checks_passed"):
        if isinstance(validation.get(key), bool):
            return "全部通过" if validation[key] else "未全部通过"
    return "以 validation.json 各项记录为准（未记录统一结论）"


def _figure(requested=None):
    if requested:
        path = Path(requested)
        path = path if path.is_absolute() else ROOT / path
    else:
        candidates = sorted((ROOT / "outputs").glob("*.png"))
        if len(candidates) != 1:
            raise RuntimeError("outputs 中须有唯一六面板 PNG，或用 --figure 指定")
        path = candidates[0]
    if not path.is_file() or not path.resolve().is_relative_to(ROOT.resolve()):
        raise RuntimeError("请先生成位于本作业目录内的结果图")
    return path


def _small(value):
    if value is None:
        return "未记录"
    value = float(value)
    return f"{value:.3e}" if 0 < abs(value) < 1e-4 else f"{value:.6f}"


def _robustness_section(document):
    path = ROOT / "outputs/robustness.json"
    if not path.is_file():
        return
    result = _json(path)
    document.add_heading("补充：测量噪声与零参考保持", 2)
    noise = result.get("noise_sensitivity", {})
    runs = noise.get("beta_runs", [])
    if runs:
        document.add_paragraph(
            f"补充噪声试验采用标准差 {_number(noise.get('std_mm_s'), 3)} mm/s、"
            f"seed={noise.get('seed', '未记录')} 的共同加性测量噪声；"
            f"重建噪声序列一致性检查：{_number(noise.get('shared_additive_noise_verified'))}。"
            "这些结果独立保存，未改变无噪声主图或平滑实际速度。")
        table = _table(document, ("β / V⁻¹", "有效信号零平台 / ms", "高通峰峰值 / V", "全程 RMSE / mm/s"))
        for item in sorted(runs, key=lambda item: item["beta"]):
            metrics = item["metrics"]
            values = (item["beta"], metrics.get("longest_flow_closed_ms"),
                      metrics.get("chatter_voltage_pp_v"), metrics.get("velocity_rmse_mm_s"))
            for cell, value, places in zip(table.add_row().cells, values, (0, 3, 5, 5)):
                cell.text = _number(value, places)
    hold = result.get("zero_reference_hold", {})
    normal, counter = hold.get("nominal_feedback_runs", {}), hold.get("feedback_disabled_counterfactual_runs", {})
    if normal and counter:
        counter_label = "关闭全部反馈"
        document.add_paragraph(
            f"保持试验从各策略正式轨迹 {hold.get('start_s', '未记录')} s 的实际阀位与速度继续，"
            f"保持零速度参考 {hold.get('duration_s', '未记录')} s，统计末 "
            f"{float(hold.get('tail_s', .1))*1000:g} ms。"
            "控制器的完整误差队列及两级低通记忆由正式测量重建，不添加速度种子或噪声。"
            f"同一策略的反事实沿用相同真实起点，{counter_label}，检验反复翻转是否由反馈维持。")
        table = _table(document, ("策略", "原反馈速度\n峰峰值 mm/s", "原反馈电压\n峰峰值 V",
                                  f"{counter_label}速度\n峰峰值 mm/s", "有效估计频率\n/ Hz"))
        for kind in KINDS:
            if kind not in normal or kind not in counter:
                continue
            metrics, counter_metrics = normal[kind]["metrics"], counter[kind]["metrics"]
            frequency = metrics.get("frequency_from_median_upcrossing_period_hz")
            frequency_text = (_number(frequency, 2) if metrics.get("frequency_estimate_meaningful")
                              and frequency is not None else "不报告")
            values = (LABELS[kind], _small(metrics.get("tail_velocity_pp_mm_s")),
                      _small(metrics.get("tail_voltage_pp_v")),
                      _small(counter_metrics.get("tail_velocity_pp_mm_s")), frequency_text)
            for cell, text in zip(table.add_row().cells, values):
                cell.text = text
        hard = normal.get("hard_over", {}).get("metrics", {})
        if hard:
            if hard.get("preceding_tail_velocity_pp_mm_s") is not None:
                document.add_paragraph(
                    f"硬过补偿前一等长窗口速度峰峰值为 {_small(hard['preceding_tail_velocity_pp_mm_s'])} mm/s，"
                    f"末段/前段幅值比为 {_number(hard.get('tail_to_preceding_pp_ratio'), 4)}。"
                "以相邻窗口幅值比较区分持续小幅运动与单次换向后逐渐衰减的振铃。")
            document.add_paragraph(
                f"硬过补偿末段记录速度过零 {hard.get('tail_speed_zero_crossings', '未记录')} 次、"
                f"完整同向上穿周期 {hard.get('tail_complete_upcrossing_cycles', '未记录')} 个。"
                "频率仅在速度幅度有诊断意义且至少存在两次同向上穿时估计；"
                "采用同向插值过零周期的中位数，而不是把单次脉冲频谱当振荡频率。"
                "该保持试验与无反馈反事实共同检验本采样等效模型的小幅极限环，"
                "不能推断任意液压对象的固有模态或全局稳定性。")
    beta_hold = hold.get('beta_hold_runs', {})
    if beta_hold:
        table = _table(document, ('β', '零参考末段速度峰峰值 / mm/s', '末段振荡频率 / Hz'))
        for key, run in beta_hold.items():
            m = run['metrics']
            for cell,value in zip(table.add_row().cells,(key,_small(m['tail_velocity_pp_mm_s']),
                    _number(m['frequency_from_median_upcrossing_period_hz'],2) if m['frequency_estimate_meaningful'] else '不报告')):
                cell.text = str(value)
    if any(key in result for key in ("passed", "all_passed", "all_checks_passed")):
        document.add_paragraph("补充试验检查状态：" + _status(result) + "；条件及原始记录见 outputs/robustness.json。")


def write_report(summary, validation, figure):
    c, strategies = summary['config'], summary['strategies']
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.65)
        section.left_margin = section.right_margin = Inches(.65)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=document.styles[name];style.font.name='Noto Sans CJK SC'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
    document.styles['Normal'].font.size=Pt(10)
    document.add_heading('第六章作业 6-3：死区逆补偿鲁棒性边界',0)
    document.add_paragraph('三策略共同模型：欠补偿平台、硬过补偿颤振与tanh光滑换向。')
    document.add_heading('1. 共同参数与试验条件',1)
    document.add_paragraph('三组使用相同的输入死区、动态有效阀口和机械速度课堂约化模型。共同的二级动态残差支路使硬逆映射的换向敏感性可在采样闭环中检验；三组只改变死区估计或补偿函数。曲线来自实际状态推进，没有注入预设振荡。')
    table=_table(document,('参数','共同设置'))
    rows=[('正弦参考',f"{c['amplitude_v']:g} V，{c['frequency_hz']:g} Hz"),
          ('采样/时长',f"{c['sample_time_s']*1000:g} ms / {c['duration_s']:g} s"),
          ('真实死区/估计',f"±{c['physical_deadzone_v']:g} V；欠{c['under_estimate_v']:g} / 过{c['over_estimate_v']:g} V"),
          ('速度增益',f"{c['gain_m_s_v']:g} m/(s·V)，5 V对应50 mm/s标度"),
          ('阀/速度时间常数',f"{c['tau_valve_s']*1000:g} / {c['tau_velocity_s']*1000:g} ms"),
          ('普通PID',f"Kp={c['kp']:g}，Ki={c['ki']:g}，Kd={c['kd']:g}"),
          ('动态残差增益/滤波',f"Kw={c['washout_gain']:g}，两级τw={c['washout_tau_s']*1000:g} ms"),
          ('整个残差缓存延迟',f"{c['feedback_delay_samples']}采样，前馈不延迟"),
          ('tanh beta',str(c['beta_values'])),('主图测量噪声',f"{c['measurement_noise_std_mm_s']:g} mm/s")]
    for label,value in rows:
        cells=table.add_row().cells;cells[0].text=label;cells[1].text=value
    document.add_paragraph('Kv是图示速度标度；阀及速度时间常数沿用给定参考。动态残差滤波、增益和缓存延迟是明确新增的共同压力测试条件，不是PPT已经给定的辨识参数。三组只改变死区估计或补偿函数，均使用相同对象、反馈、初态及采样。')
    document.add_heading('2. 对象和控制律',1)
    document.add_paragraph('u_eff=Dδ(u)=sign(u)·max(|u|−δ,0)；τv·dz/dt=u_eff−z；τm·dv/dt=Kv·z−v。u_eff为题图静态死区输出，z为滤波后的有效阀口等效电压，二者不能混用；u_eff=0不代表动态阀口立即无流量。两级动态在每个ZOH输入区间解析积分，速度没有硬置零或后处理平滑。')
    document.add_paragraph('e[k]=(vd[k]−v_meas[k])/Kv；ed=e[k−d]；a=exp(−Ts/τw)。按采样顺序：L1[k]=aL1[k−1]+(1−a)ed；h1=ed−L1[k]；L2[k]=aL2[k−1]+(1−a)h1；h2=h1−L2[k]。u_lin=(vd+τm·ad)/Kv+Kw·h2。两级 washout 的 DC 增益为零，故保留慢变死区估计失配；它对动态残差和采样延迟敏感，用来测试硬逆函数的鲁棒性。')
    document.add_paragraph('硬补偿u=u_lin+δhat·sign(u_lin)，|u_lin|≤1e−10 V时输出零；光滑补偿u=u_lin+δhat·tanh(βu_lin)。最终执行电压限于±10 V。该动态残差测试架构不是原通用模块三的完整双腔逆动力学PID；质量、压力、泄漏和Stribeck摩擦未在本降阶通道中计算。原模块三保存在materials供比较。')
    document.add_heading('3. 六联图与定量结果',1)
    document.add_paragraph(f"换向窗口{c['zoom_start_s']:g}–{c['zoom_end_s']:g}s。有效信号零平台通过|u|≤δ的采样阈值重建统计；近零速度另外按|v|≤{c['velocity_flat_threshold_mm_s']:g}mm/s统计。两者不等同。给定5 V、0.1 Hz和0.10 V残余静态死区时平台约为63.67 ms；示意图标注的约110 ms不作为本次实测值。")
    document.add_paragraph(f"电压和速度颤振采用4阶{c['chatter_cutoff_hz']:g}Hz零相位高通分量统计，滤波只定义指标，不进入控制环或所绘曲线。单次快速换向也含高频分量，持续极限环需用零参考保持和停反馈反事实检验。")
    table=_table(document,('指标',*[LABELS[k] for k in KINDS]))
    rows=[('longest_effective_zero_ms','有效信号最长零平台 / ms',3),
          ('longest_velocity_near_zero_ms','近零速度最长平台 / ms',3),
          ('chatter_voltage_pp_v','电压高频峰峰值 / V',4),
          ('chatter_velocity_pp_mm_s','速度高频峰峰值 / mm/s',4),
          ('velocity_excess_variation_mm_s','速度超额变差 / mm/s',4),
          ('linear_sign_reversals','u_lin符号翻转次数',0),
          ('velocity_rmse_mm_s','全周期速度RMSE / mm/s',5),
          ('max_abs_voltage_v','最大执行电压绝对值 / V',4),
          ('sampled_saturation_count','电压饱和采样数',0)]
    for key,label,digits in rows:
        cells=table.add_row().cells;cells[0].text=label
        for cell,kind in zip(cells[1:],KINDS):cell.text=_number(strategies[kind]['metrics'][key],digits)
    sec=document.sections[0]
    document.add_picture(str(figure),width=sec.page_width-sec.left_margin-sec.right_margin)
    document.add_paragraph('图 1：直接绘制原始闭环记录；左列为全周期 (a1–a3)，右列为换向局部 (b1–b3)。图中有效信号为 D(u)，动态阀口状态另存 CSV。两组硬符号逆映射在零点均不连续。')
    document.add_heading('4. beta扫描和数学讨论',1)
    table=_table(document,('beta / V^-1','有效信号零平台/ms','近零速度/ms','电压高频pp/V','速度高频pp/mm/s','RMSE/mm/s'))
    for item in summary['beta_scan']:
        m=item['metrics'];values=(item['beta'],m['longest_effective_zero_ms'],m['longest_velocity_near_zero_ms'],m['chatter_voltage_pp_v'],m['chatter_velocity_pp_mm_s'],m['velocity_rmse_mm_s'])
        for cell,value,digits in zip(table.add_row().cells,values,(0,3,3,4,4,5)):cell.text=_number(value,digits)
    document.add_paragraph('欠补偿静态残留0.10 V区间；5 V、0.1 Hz正弦的静态平台为2asin(0.10/5)/(2π·0.1)=63.67ms，机械近零速度平台可更宽。硬过补偿在零点两侧的命令极限为±0.60 V，死区输出极限为±0.10 V；反馈可维持反复换向。仅一次sign跳变不等于任意对象必然失稳。')
    document.add_paragraph('任何有限beta下，g(q)=q+0.6tanh(beta*q)连续且g(0)=0，必有正阈值g(q0)=0.5。因此光滑补偿仍有残余平台；完整映射零点斜率为1+0.6beta。增大beta通常缩短平台，但增强动态残差敏感性。主换向窗口与零参考保持的结论须分别报告，不能把所有tanh参数都称为无振荡。')
    _robustness_section(document)
    document.add_heading('5. 验证与复现',1)
    document.add_paragraph('本次数据与机制检查：'+_status(validation)+f"，共 {len(validation['checks'])} 项。"
                           '另有 13 项单元测试，日志见 outputs/unit_tests.txt。输入死区对象的解析推进以独立高精度 ODE 和半步等价验证；误差缓存、滤波零 DC 增益、保持段记忆连续及完整停反馈均单独检查。')
    document.add_paragraph('运行：python -m pip install -r requirements.txt；python work_6_3.py；python -m unittest -v test_work_6_3.py；python package_delivery.py。CSV状态记录在当前指令作用前，u及D(u)作用于下一个Ts区间。全部代码、配置、PPT、原始CSV、报告和唯一PNG均在ZIP内，manifest列SHA256。')
    document.add_paragraph('图中三类失配现象均由公开参数下的课堂等效模型计算，示意图数值及波纹周期不作为硬编码目标。该模型不是完整双腔数字孪生或设备辨识结论；压力安全和全局稳定性不在本实验验证范围内。')
    path=ROOT/'deliverables'/REPORT_NAME;path.parent.mkdir(parents=True,exist_ok=True);document.save(path)
    return path


def included_files():
    return sorted((path for path in ROOT.rglob("*") if path.is_file()
                   and not EXCLUDED_DIRS.intersection(path.relative_to(ROOT).parts)
                   and path.name != "manifest.json"
                   and path.suffix.lower() not in {".zip", ".pyc", ".pyo", ".tmp"}),
                  key=lambda path: path.relative_to(ROOT).as_posix())


def package(figure=None):
    required = {"README.md", "requirements.txt", "experiment_config.json", "work_6_3.py", "simulation.py",
                "deadzone_model.py", "input_deadzone_model.py", "model_assisted_controller.py", "plotting.py", "package_delivery.py",
                "materials/chap06-ppt.pdf", "materials/module3_reference.py", "outputs/summary.json",
                "outputs/validation.json", "outputs/under.csv", "outputs/hard_over.csv", "outputs/smooth.csv",
                "outputs/beta_5.csv", "outputs/beta_18.csv", "outputs/beta_50.csv", "outputs/robustness.json"}
    available = {path.relative_to(ROOT).as_posix() for path in included_files()}
    if not required <= available:
        raise RuntimeError(f"交付文件尚未齐全：{sorted(required-available)}")
    figure = _figure(figure)
    write_report(_json(ROOT / "outputs/summary.json"), _json(ROOT / "outputs/validation.json"), figure)
    files = included_files()
    manifest = {"project": "homework6_3", "files": {
        path.relative_to(ROOT).as_posix(): {"bytes": path.stat().st_size,
                                        "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
        for path in files}}
    manifest_path = ROOT / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    destination = ROOT / "deliverables" / ARCHIVE_NAME
    with ZipFile(destination, "w", ZIP_DEFLATED, compresslevel=9) as archive:
        for path in [*files, manifest_path]:
            archive.write(path, (Path(ROOT.name) / path.relative_to(ROOT)).as_posix())
    with ZipFile(destination) as archive:
        broken = archive.testzip()
        if broken is not None:
            raise RuntimeError(f"ZIP 完整性检查失败：{broken}")
        for relative, entry in manifest["files"].items():
            payload = archive.read(f"{ROOT.name}/{relative}")
            if len(payload) != entry["bytes"] or hashlib.sha256(payload).hexdigest() != entry["sha256"]:
                raise RuntimeError(f"ZIP 文件与清单不符：{relative}")
    print(f"报告：{ROOT / 'deliverables' / REPORT_NAME}")
    print(f"ZIP：{destination}（{len(files)+1} 个文件，{destination.stat().st_size} 字节）")
    return destination


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--figure", help="指定本作业目录内已生成的六面板 PNG")
    args = parser.parse_args()
    package(args.figure)


if __name__ == "__main__":
    main()
