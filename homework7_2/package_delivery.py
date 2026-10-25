#!/usr/bin/env python3
"""由习题 7-2 的真实运行结果生成中文实验报告、校验清单与交付 ZIP。"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
REPORT = '第七章_习题7-2_预测时域实验报告.docx'
ARCHIVE = '第七章_习题7-2_预测时域完整交付.zip'
EXCLUDED_PARTS = {'__pycache__', '.git', '.pytest_cache', '.mypy_cache', '.ruff_cache', '.cache'}


def _read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def _fmt(value, precision=3):
    return '未记录' if value is None else f'{float(value):.{precision}f}'


def _field(metrics, *path):
    item = metrics
    for key in path:
        if not isinstance(item, dict):
            return None
        item = item.get(key)
    return item


def _style(document):
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(.62)
    section.left_margin = section.right_margin = Inches(.68)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.name = 'Noto Sans CJK SC'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    document.styles['Normal'].font.size = Pt(9.5)
    document.styles['Title'].font.size = Pt(18)


def _table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = 'Light Shading Accent 1'
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = str(label)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = str(value)
    return table


def build_report(summary, validation, figure, destination):
    """Build a reviewable report from recorded metrics, never reference numbers."""
    records = summary.get('results', {})
    config = summary.get('config', {})
    horizons = [int(n) for n in config.get('horizons', [5, 10, 15, 20, 30, 40, 50])]
    missing = [n for n in horizons if str(n) not in records]
    if missing:
        raise RuntimeError(f'summary.json 缺少预测时域 {missing} 的结果')
    figure = Path(figure)
    if not figure.is_file():
        raise FileNotFoundError(f'四联图不存在：{figure}')
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    report = Document()
    _style(report)
    report.add_heading('第七章习题 7-2：液压 MPC 预测时域与安全—算力权衡', 0)
    report.add_paragraph('基于七组真实闭环轨迹、CVXPY 滚动优化与高密压力记录的可复现实验')

    report.add_heading('1. 实验范围与共同条件', 1)
    report.add_paragraph(
        f"同一非对称阀控液压缸及相同初态下，于 t={_fmt(config.get('step_time'), 3)} s "
        f"输入 {_fmt(config.get('target_mm'), 1)} mm 阶跃指令；采样周期 "
        f"Ts={_fmt(1000*config['ts'], 1)} ms，仿真至 {_fmt(config.get('duration'), 3)} s。"
        f"扫描 N={horizons}，分别代表 N·Ts 的预测视野。"
        f"输入约束 |u|≤10 V、|Δu|≤{_fmt(config.get('max_du_step'), 1)} V/步；"
        f"有杆腔实际压力安全线为 {config['p_safe']:g} MPa，预测节点采用"
        f"{config['p_safe']-config['pressure_margin']:g} MPa 上限（含"
        f"{config['pressure_margin']:g} MPa 裕量）。这些参数是公开课堂仿真设置，"
        "并非由实机数据辨识得到。")
    report.add_paragraph(
        '真实对象为带硬死区、静摩擦事件、有向孔口、泄漏与阀动态的混合液压模型；'
        '控制器使用可微灰箱近似。每个采样时刻以真实对象状态重置预测，只执行最优序列首个电压。'
        '压力峰值从真实积分器的高密轨迹读取，包含采样间峰值；未削顶或平滑真实压力。')
    if config.get('plant_config_file'):
        report.add_paragraph(
            f"本次实验通过 {config['plant_config_file']} 启用显式教学等效回路："
            '伸出供油口与回油旁路分别设定流量能力，阀芯和活塞近静止时开放双腔压力调节支路。'
            '支路流量写入真实对象及预测模型的压力微分方程，不重置、平移或裁剪压力记录。'
            '此回路及参数依据参考图形提出，属于公开的课堂假设；单张参考图无法辨识真实装置的回路。'
            '移除 experiment_config.json 中的 plant_config_file 可运行原闭中位对称计量口对象。')

    report.add_heading('2. 模型与优化问题', 1)
    report.add_paragraph(
        '物理状态 z=[相对位移 mm，速度 m/s，P1 MPa，P2 MPa，阀位等效电压 V] 为五维；'
        '将上一时刻实际电压 u(k−1) 纳入后，增广状态为六维。课件中的“五维增广”'
        '对应忽略阀动态的一种简化模型；本实验保留真实阀惯性，所以矩阵尺寸不同。'
        '逐步雅可比线性化后离散为 z(k+1)=A(k)z(k)+B(k)u(k)+c(k)，'
        '堆叠预测写为 X=F·x_aug(k)+Φ·ΔU+b。')
    report.add_paragraph(
        '以位置/速度误差和电压增量构造凸二次目标，并将电压、电压增量及 P2 预测节点上界'
        '写成 A_qp·ΔU 的线性不等式。CVXPY 缓存满足 DPP 的问题结构，底层调用 OSQP；'
        'QP 解通过原始可行性和 KKT 残差校核后，仍需用非线性灰箱重新验证未来输入与压力。'
        '若局部 QP 无可用计划，程序可调用公开的有限网格恢复策略：测试投影后的候选控制序列，'
        '只接受经过同一非线性预测约束验证的序列，且把该步骤单独计数。此举不是纯 QP 求解成功，'
        '也不构成一般的递归可行性证明。')
    if config.get('terminal_tail_weight0', 0.) > 0.:
        report.add_paragraph(
            f"所有时域共用终端制动代价：终端速度权重乘以 "
            f"{_fmt(config.get('terminal_velocity_factor', 1.), 1)}；"
            f"以 {_fmt(1000*config.get('terminal_stop_time_s', 0.), 1)} ms "
            '估计视野外的惯性滑行，平滑地惩罚预计停止位置超过目标。'
            f"越程代价权重采用 {config['terminal_tail_weight0']:g}·"
            f"exp[−N·Ts/{_fmt(1000*config['terminal_tail_decay_s'], 1)} ms]，"
            '对全部 N 使用同一公式；局部 QP 求其二次近似，非线性候选按原代价重新评分。'
            '这是控制器的终端预测假设，不代表新增物理制动器。')
    report.add_paragraph(
        '另外，当 CVXPY 局部 QP 没有产出新的非线性可行计划、控制器仅保留已经重新验证的'
        '历史移位计划时，程序可对该移位计划进行有限网格改进：原计划仍作为可行候选，'
        '只在另一候选满足相同压力与输入约束、且代价改善超过阈值时替换。逐步记录状态码 4 '
        '和改进次数；它是已验证计划的附加候选搜索，不是纯最优 QP 解，也不提供一般递归可行性证明。')
    report.add_paragraph(
        '本实验中 Φ 尺寸为 (6N)×N，海森矩阵 H 为 N×N，矩阵条目随 N 按 O(N²) 增长；'
        '密集分解的典型代价可达 O(N³)，并非课题描述中的“指数级维度膨胀”。'
        '实际耗时还包含线性化、非线性预测、CVXPY 调用、系统调度与硬件环境。')

    report.add_heading('3. 七组实际测量结果', 1)
    _table(report,
           ('N', '视野 ms', '位移超调 %', '真实 P2 峰 MPa', '单次 QP 中位 ms',
            'CVXPY 均值 ms', '完整控制中位 ms', '失败恢复次数', '移位改进次数', '状态'),
           [(n, _fmt(_field(records[str(n)], 'prediction_ms'), 0),
             _fmt(_field(records[str(n)], 'overshoot_pct'), 3),
             _fmt(_field(records[str(n)], 'peak_pressure_mpa'), 4),
             _fmt(_field(records[str(n)], 'solver_attempt', 'median_ms'), 3),
             _fmt(_field(records[str(n)], 'cvxpy_call', 'mean_ms'), 3),
             _fmt(_field(records[str(n)], 'whole_controller', 'median_ms'), 3),
             records[str(n)].get('restoration_count', 0),
             records[str(n)].get('refinement_count', 0),
             '完成' if records[str(n)].get('status') == 'completed' else '未完成')
            for n in horizons])
    report.add_paragraph(
        '单次 QP 中位数逐次统计 OSQP 底层求解时间，与图 (d) 口径一致；'
        'CVXPY 栏为每控制周期内 Problem.solve 的墙钟时间合计均值；'
        '完整控制时间还含模型预测和非线性验证。三者不能互换。'
        '所有数值保存在逐步 CSV 和求解日志中。')
    deadline_rows = []
    for n in horizons:
        item = records[str(n)]
        fraction = _field(item, 'whole_controller', 'over_1ms_fraction')
        deadline_rows.append((n, _fmt(_field(item, 'solver_attempt', 'p95_ms'), 3),
                              _fmt(_field(item, 'solver', 'mean_ms'), 3),
                              _fmt(_field(item, 'cvxpy_call', 'p95_ms'), 3),
                              _fmt(_field(item, 'whole_controller', 'p95_ms'), 3),
                              _fmt(100*fraction, 1) if fraction is not None else '未记录',
                              _fmt(item.get('pressure_violation_mpa'), 5)))
    _table(report,
           ('N', '单次 QP P95 ms', '周期 QP 均值 ms', '周期 CVXPY P95 ms', '完整控制 P95 ms',
            '完整控制 >1ms 占比 %', '真实 P2 越限 MPa'), deadline_rows)
    report.add_paragraph(
        '一周期内可能连续求解多个局部 QP，所以周期 QP 均值可能明显高于单次中位数。'
        '采样周期 1 ms 是理论实时截止线；本表是在当前通用计算环境的观测时间，'
        '不能代替实时操作系统、指定处理器和最坏情况执行时间认证。')
    deadline_failures = [n for n in horizons
                         if (_field(records[str(n)], 'whole_controller', 'over_1ms_fraction')
                             or 0.) > 0.]
    if deadline_failures:
        report.add_paragraph(
            f'本次完整控制步骤在 N={deadline_failures} 出现超过 1 ms 的周期；'
            '不能凭底层求解器的部分低耗时样本宣称整套 Python/CVXPY 控制器满足'
            '1 ms 硬实时要求，也不能据此直接认定课件示意的 N∈[15,25] 已通过算力验收。'
            '工程整定应在目标处理器上复测最坏情况执行时间。')

    report.add_heading('4. 观察与边界解释', 1)
    for n, name in ((5, '短视'), (20, '基准'), (50, '长视')):
        item = records[str(n)]
        status = '完成' if item.get('status') == 'completed' else f"于 {item.get('final_time_s')} s 提前终止"
        report.add_paragraph(
            f"N={n}（{name}，{n} ms）{status}：位移超调 "
            f"{_fmt(item.get('overshoot_pct'), 3)}%，高密真实 P2 峰值 "
            f"{_fmt(item.get('peak_pressure_mpa'), 4)} MPa；有限网格压力可行恢复 "
            f"{item.get('restoration_count', 0)} 次；经验证的移位计划再经有限网格改进 "
            f"{item.get('refinement_count', 0)} 次。底层 OSQP 平均 "
            f"{_fmt(_field(item, 'solver', 'mean_ms'), 3)} ms，CVXPY 调用平均 "
            f"{_fmt(_field(item, 'cvxpy_call', 'mean_ms'), 3)} ms，完整控制中位 "
            f"{_fmt(_field(item, 'whole_controller', 'median_ms'), 3)} ms。")
    if any(_field(records[str(n)], 'overshoot_pct') > _field(records[str(n+5)], 'overshoot_pct')
           for n in (5, 10) if str(n+5) in records):
        report.add_paragraph(
            '超调并不必然随预测时域单调下降。逐步非线性局部优化、控制权重切换和'
            '恢复策略都可能造成局部差异；应根据七组原始曲线而不是示意图推断。')
    all_safe = all(_field(records[str(n)], 'peak_pressure_mpa') <= config['p_safe']
                   for n in horizons)
    all_complete = all(records[str(n)].get('status') == 'completed' for n in horizons)
    if all_safe and all_complete:
        safety_text = '本次七组完整闭环的高密真实压力峰值均未超过 21 MPa。'
    elif all_safe:
        safety_text = '已观测的高密真实压力峰值未超过 21 MPa；提前终止组的后续安全性未知。'
    else:
        safety_text = '本次至少有一组高密真实压力峰值超过 21 MPa，具体越限见上表。'
    report.add_paragraph(
        safety_text
        + '有限仿真中守住安全线，不等于任意模型误差、扰动与实物系统均安全。'
        '当前硬约束针对 P2 上界，未保护 P1 上界或两腔压力下界。')
    if config.get('plant_config_file'):
        report.add_paragraph(
            '原闭中位对称计量口对象的巡航回油压差和制动后封闭腔压，'
            '与课堂示意图的低巡航压力及约 2 MPa 尾压差异较大。'
            '本次教学等效回路显式加入回油旁路与近静止压力调节支路来研究这一差异；'
            '图和表中的压力仍来自真实积分轨迹，未修改记录。'
            '该结果仅说明所声明回路假设下的响应，不能由图形反推出真实设备具有同一结构。')
    else:
        report.add_paragraph(
            '本次采用原闭中位对称计量口对象。巡航时有杆腔需形成回油压差；'
            '阀芯回中位后两腔封闭，外泄漏系数为零，制动后压力不会自动回到初始基线。'
            '图和表中的压力来自真实积分轨迹，未修改记录。')
    report.add_paragraph(
        f"验证脚本总体结果：{'通过' if validation.get('passed') else '未全部通过'}。"
        '验证内容包含图表数值与 CSV、真实积分峰值、输入约束及预测压力节点的独立核对；'
        '如开启严格复算，见 outputs/validation.json 的 recorded-input replay 数据。')

    report.add_heading('5. 四联图与可复现文件', 1)
    report.add_picture(str(figure), width=Inches(10.0))
    report.add_paragraph(
        '图 (a)–(c) 使用 N=5/20/50 的闭环样本、实际控制电压和高密 P2；'
        '图 (d) 绘制七组单次 QP 耗时中位数、均值与真实压力峰值。'
        '运行命令：python work_7_2.py；随后 python verify_experiment.py --strict；'
        '再运行 python package_delivery.py。源代码、配置、原始数据、逐步诊断、'
        '课件资料、reference_comparison.md 差异诊断和 SHA256 清单均随 ZIP 一同提供。')
    report.save(destination)
    return destination


def _included_files(root, archive_path):
    for path in sorted(root.rglob('*')):
        if not path.is_file() or any(part in EXCLUDED_PARTS for part in path.relative_to(root).parts):
            continue
        if path.resolve() == archive_path.resolve() or path.name == 'manifest.json':
            continue
        if path.suffix in ('.pyc', '.pyo'):
            continue
        yield path


def build_package(root=ROOT, output_dir=None, deliverables_dir=None, figure=None):
    root = Path(root).resolve()
    output_dir = Path(output_dir) if output_dir is not None else root/'outputs'
    deliverables_dir = Path(deliverables_dir) if deliverables_dir is not None else root/'deliverables'
    if not output_dir.is_absolute():
        output_dir = root/output_dir
    if not deliverables_dir.is_absolute():
        deliverables_dir = root/deliverables_dir
    if not output_dir.resolve().is_relative_to(root) or not deliverables_dir.resolve().is_relative_to(root):
        raise ValueError('输出与交付目录必须位于本作业文件夹内')
    summary_path = output_dir/'summary.json'
    validation_path = output_dir/'validation.json'
    if not summary_path.is_file() or not validation_path.is_file():
        raise FileNotFoundError('请先运行完整时域扫描与 verify_experiment.py，生成 summary.json / validation.json')
    figure = Path(figure) if figure is not None else output_dir/'horizon_comparison.png'
    if not figure.is_absolute():
        figure = root/figure
    if not figure.resolve().is_relative_to(root):
        raise ValueError('四联图必须位于作业文件夹内')
    summary, validation = _read_json(summary_path), _read_json(validation_path)
    deliverables_dir.mkdir(parents=True, exist_ok=True)
    report_path = build_report(summary, validation, figure, deliverables_dir/REPORT)
    archive_path = deliverables_dir/ARCHIVE
    files = list(_included_files(root, archive_path))
    manifest = dict(
        title='第七章习题 7-2：MPC 预测时域与安全—算力权衡',
        root_folder=root.name,
        archive_name=ARCHIVE,
        figure=str(figure.relative_to(root)),
        report=str(report_path.relative_to(root)),
        validation_passed=bool(validation.get('passed')),
        sha256={str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
                for path in files},
    )
    manifest_path = root/'manifest.json'
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    with ZipFile(archive_path, 'w', ZIP_DEFLATED, compresslevel=8) as archive:
        for path in files + [manifest_path]:
            archive.write(path, arcname=str(Path(root.name)/path.relative_to(root)))
    with ZipFile(archive_path) as archive:
        corrupt = archive.testzip()
        if corrupt is not None:
            raise RuntimeError(f'交付 ZIP 校验失败：{corrupt}')
    return dict(report=str(report_path), archive=str(archive_path),
                manifest=str(manifest_path), files=len(files)+1)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=ROOT/'outputs')
    parser.add_argument('--deliverables-dir', type=Path, default=ROOT/'deliverables')
    parser.add_argument('--figure', type=Path)
    args = parser.parse_args()
    print(json.dumps(build_package(output_dir=args.output_dir,
                                   deliverables_dir=args.deliverables_dir,
                                   figure=args.figure), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
